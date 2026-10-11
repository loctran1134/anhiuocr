import os
import sys
import re
import json
import webbrowser
import zipfile
import tempfile
import subprocess
from datetime import datetime
from typing import List, Dict, Any

# Windows console & encoding setup
if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

# Configure local models directory for 100% offline bundled usage (Supports PyInstaller .exe)
if getattr(sys, 'frozen', False):
    candidates = [
        getattr(sys, '_MEIPASS', ''),
        os.path.join(os.path.dirname(sys.executable), '_internal'),
        os.path.dirname(sys.executable)
    ]
else:
    candidates = [os.path.dirname(os.path.abspath(__file__))]

for c in candidates:
    if c:
        m_dir = os.path.join(c, 'models')
        if os.path.exists(m_dir):
            os.environ["PADDLE_PDX_CACHE_HOME"] = m_dir
            break

# Windows DLL directory fix for PyTorch / Torchvision / Paddle
try:
    import torch
    torch_lib = os.path.join(os.path.dirname(torch.__file__), 'lib')
    if os.path.exists(torch_lib):
        try:
            os.add_dll_directory(torch_lib)
        except Exception:
            pass
except Exception:
    pass

try:
    import torchvision
    tv_dir = os.path.dirname(torchvision.__file__)
    if os.path.exists(tv_dir):
        try:
            os.add_dll_directory(tv_dir)
        except Exception:
            pass
except Exception:
    pass

import cv2
import numpy as np
import pandas as pd

from ocr_engine import TimemarkOCREngine

_GLOBAL_OCR_ENGINE = None

def get_ocr_engine() -> TimemarkOCREngine:
    global _GLOBAL_OCR_ENGINE
    if _GLOBAL_OCR_ENGINE is None:
        _GLOBAL_OCR_ENGINE = TimemarkOCREngine()
    return _GLOBAL_OCR_ENGINE

from PySide6.QtCore import Qt, QThread, Signal, QByteArray
from PySide6.QtGui import (
    QIcon, QPixmap, QImage, QPainter, QColor, QFont
)
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QLabel, QPushButton, QLineEdit, QTableWidget, QTableWidgetItem,
    QHeaderView, QFileDialog, QMessageBox, QScrollArea, QFrame,
    QProgressBar, QDialog, QAbstractItemView, QGridLayout, QComboBox,
    QTextEdit, QTabWidget, QProgressDialog
)
from PySide6.QtSvg import QSvgRenderer

# ----------------- VECTOR SVG ICONS (NO EMOJIS) -----------------
SVG_CAMERA = '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" fill="none" stroke="{color}" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M14.5 4h-5L7 7H4a2 2 0 0 0-2 2v9a2 2 0 0 0 2 2h16a2 2 0 0 0 2-2V9a2 2 0 0 0-2-2h-3l-2.5-3z"/><circle cx="12" cy="13" r="3"/></svg>'''

SVG_CLOUD_UPLOAD = '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" fill="none" stroke="{color}" stroke-width="1.8" stroke-linecap="round" stroke-linejoin="round"><path d="M4 14.899A7 7 0 1 1 15.71 8h1.79a4.5 4.5 0 0 1 2.5 8.242"/><path d="M12 12v9"/><path d="m16 16-4-4-4 4"/></svg>'''

SVG_IMAGE = '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" fill="none" stroke="{color}" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><rect width="18" height="18" x="3" y="3" rx="2" ry="2"/><circle cx="9" cy="9" r="2"/><path d="m21 15-3.086-3.086a2 2 0 0 0-2.828 0L6 21"/></svg>'''

SVG_CHECK_CIRCLE = '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" fill="none" stroke="{color}" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M22 11.08V12a10 10 0 1 1-5.93-9.14"/><polyline points="22 4 12 14.01 9 11.01"/></svg>'''

SVG_ALERT_TRIANGLE = '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" fill="none" stroke="{color}" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="m21.73 18-8-14a2 2 0 0 0-3.48 0l-8 14A2 2 0 0 0 4 21h16a2 2 0 0 0 1.73-3Z"/><line x1="12" y1="9" x2="12" y2="13"/><line x1="12" y1="17" x2="12.01" y2="17"/></svg>'''

SVG_MAP_PIN = '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" fill="none" stroke="{color}" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M20 10c0 6-8 12-8 12s-8-6-8-12a8 8 0 0 1 16 0Z"/><circle cx="12" cy="10" r="3"/></svg>'''

SVG_SPARKLE = '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" fill="none" stroke="{color}" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="m12 3-1.912 5.813a2 2 0 0 1-1.275 1.275L3 12l5.813 1.912a2 2 0 0 1 1.275 1.275L12 21l1.912-5.813a2 2 0 0 1 1.275-1.275L21 12l-5.813-1.912a2 2 0 0 1-1.275-1.275L12 3Z"/></svg>'''

def get_avatar_initials(name: str) -> str:
    if not name or name == "--":
        return "?"
    words = name.strip().split()
    if len(words) == 1:
        return words[0][:2].upper()
    return (words[0][0] + words[-1][0]).upper()

AVATAR_COLORS = [
    ("#eff6ff", "#1d4ed8", "#bfdbfe"),  # Blue
    ("#fdf4ff", "#a21caf", "#f5d0fe"),  # Fuchsia
    ("#f0fdf4", "#15803d", "#bbf7d0"),  # Green
    ("#fff7ed", "#c2410c", "#fed7aa"),  # Orange
    ("#faf5ff", "#7e22ce", "#e9d5ff"),  # Purple
    ("#f0fdfa", "#0f766e", "#99f6e4"),  # Teal
]

def get_avatar_color(name: str):
    h = sum(ord(c) for c in (name or ""))
    return AVATAR_COLORS[h % len(AVATAR_COLORS)]

SVG_EXCEL = '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" fill="none" stroke="{color}" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M21 15v4a2 2 0 0 1-2 2H5a2 2 0 0 1-2-2v-4"/><polyline points="7 10 12 15 17 10"/><line x1="12" x2="12" y1="15" y2="3"/></svg>'''

SVG_TRASH = '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" fill="none" stroke="{color}" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M3 6h18"/><path d="M19 6v14c0 1-1 2-2 2H7c-1 0-2-1-2-2V6"/><path d="M8 6V4c0-1 1-2 2-2h4c1 0 2 1 2 2v2"/></svg>'''

SVG_EMPTY_ILLUST = '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 120 120" fill="none">
<rect x="25" y="25" width="55" height="55" rx="8" fill="#EFF6FF" stroke="#BFDBFE" stroke-width="2.5"/>
<path d="M35 62L46 50C47.5 48.5 50 48.5 51.5 50L68 68" stroke="#93C5FD" stroke-width="2.5" stroke-linecap="round"/>
<circle cx="45" cy="42" r="5" fill="#93C5FD"/>
<circle cx="78" cy="78" r="18" fill="white" stroke="#2563EB" stroke-width="3"/>
<line x1="91" y1="91" x2="105" y2="105" stroke="#2563EB" stroke-width="4" stroke-linecap="round"/>
<line x1="72" y1="78" x2="84" y2="78" stroke="#93C5FD" stroke-width="2" stroke-linecap="round"/>
</svg>'''

def get_svg_pixmap(svg_template: str, color: str = "#2563eb", width: int = 24, height: int = 24) -> QPixmap:
    svg_str = svg_template.replace("{color}", color)
    renderer = QSvgRenderer(QByteArray(svg_str.encode('utf-8')))
    pix = QPixmap(width, height)
    pix.fill(Qt.transparent)
    painter = QPainter(pix)
    renderer.render(painter)
    painter.end()
    return pix

def get_svg_icon(svg_template: str, color: str = "#2563eb", size: int = 20) -> QIcon:
    return QIcon(get_svg_pixmap(svg_template, color, size, size))

def safe_load_pixmap(file_path: str, max_w: int = 80, max_h: int = 80) -> QPixmap:
    try:
        with open(file_path, 'rb') as f:
            data = f.read()
        qimg = QImage()
        qimg.loadFromData(data)
        if qimg.isNull():
            return QPixmap()
        return QPixmap.fromImage(qimg).scaled(max_w, max_h, Qt.KeepAspectRatio, Qt.SmoothTransformation)
    except Exception:
        return QPixmap()

def cv2_to_qpixmap(cv_img: np.ndarray, max_w: int = None, max_h: int = None) -> QPixmap:
    try:
        if cv_img is None or cv_img.size == 0:
            return QPixmap()
        h, w = cv_img.shape[:2]
        rgb = cv2.cvtColor(cv_img, cv2.COLOR_BGR2RGB)
        bytes_per_line = 3 * w
        qimg = QImage(rgb.data, w, h, bytes_per_line, QImage.Format_RGB888)
        pix = QPixmap.fromImage(qimg)
        if max_w and max_h:
            return pix.scaled(max_w, max_h, Qt.KeepAspectRatio, Qt.SmoothTransformation)
        return pix
    except Exception:
        return QPixmap()

def safe_read_cv2(file_path: str):
    try:
        from PIL import Image, ImageOps
        pil_img = Image.open(file_path)
        pil_img = ImageOps.exif_transpose(pil_img)
        return cv2.cvtColor(np.array(pil_img.convert("RGB")), cv2.COLOR_RGB2BGR)
    except Exception:
        pass
    try:
        with open(file_path, 'rb') as f:
            data = f.read()
        nparr = np.frombuffer(data, np.uint8)
        return cv2.imdecode(nparr, cv2.IMREAD_COLOR)
    except Exception:
        return None

def format_clean_address(addr: str) -> str:
    if not addr or "không tìm thấy" in addr.lower() or "không xác định" in addr.lower():
        return addr if addr else "Chưa xác định địa chỉ"
    cleaned = re.sub(r'[\,\s]*\d{1,2}\.\d{4,8}[\s\?°]*[nN]?[\,\s]+\d{2,3}\.\d{4,8}[\s\?°\-]*[eE]?', '', addr)
    cleaned = re.sub(r'^(?:xã\s+)?địa\s*chỉ\s*[:\-\s]*', '', cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r'^(?:đ\/c|address|location)\s*[:\-\s]*', '', cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r'^[a-zA-Z0-9][\.\,]\s*', '', cleaned.strip())
    parts = [p.strip() for p in cleaned.split(',')]
    formatted = []
    for p in parts:
        if p:
            words = p.split()
            formatted.append(" ".join(w.capitalize() for w in words))
    return ", ".join(formatted) if formatted else cleaned


# ----------------- OCR ENGINE PRELOADER -----------------
class EnginePreloadThread(QThread):
    def run(self):
        try:
            get_ocr_engine()
        except Exception:
            pass


# ----------------- OCR WORKER THREAD -----------------
class OCRWorkerThread(QThread):
    progress_update = Signal(dict)
    finished_success = Signal(dict)
    finished_error = Signal(str)

    def __init__(self, file_paths: List[str]):
        super().__init__()
        self.file_paths = file_paths

    def run(self):
        try:
            engine = get_ocr_engine()

            image_items = []
            for path in self.file_paths:
                img = safe_read_cv2(path)
                if img is not None:
                    image_items.append({
                        "filename": os.path.basename(path),
                        "filepath": path,
                        "image": img
                    })

            if not image_items:
                self.finished_error.emit("Không thể tải các file ảnh đã chọn!")
                return

            def on_progress(p_data):
                self.progress_update.emit(p_data)

            res = engine.process_images(image_items, progress_callback=on_progress)
            self.finished_success.emit(res)
        except Exception as e:
            import traceback
            traceback.print_exc()
            self.finished_error.emit(str(e))


# ----------------- THUMBNAIL CARD WIDGET -----------------
class ThumbnailCard(QFrame):
    delete_clicked = Signal(str)

    def __init__(self, file_path: str, parent=None):
        super().__init__(parent)
        self.file_path = file_path
        self.setFixedSize(60, 60)
        self.setStyleSheet("""
            ThumbnailCard {
                background-color: #f1f5f9;
                border: 1px solid #cbd5e1;
                border-radius: 6px;
            }
        """)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(2, 2, 2, 2)

        self.img_label = QLabel()
        self.img_label.setAlignment(Qt.AlignCenter)
        self.img_label.setStyleSheet("border: none; background: transparent;")
        pix = safe_load_pixmap(file_path, 54, 54)
        if not pix.isNull():
            self.img_label.setPixmap(pix)
        else:
            self.img_label.setText("IMG")
        layout.addWidget(self.img_label)

        self.close_btn = QPushButton("×", self)
        self.close_btn.setFixedSize(16, 16)
        self.close_btn.move(42, 2)
        self.close_btn.setCursor(Qt.PointingHandCursor)
        self.close_btn.setStyleSheet("""
            QPushButton {
                background-color: rgba(15, 23, 42, 0.75);
                color: #ffffff;
                font-size: 12px;
                font-weight: bold;
                border: none;
                border-radius: 8px;
                padding-bottom: 2px;
            }
            QPushButton:hover {
                background-color: #ef4444;
            }
        """)
        self.close_btn.clicked.connect(lambda: self.delete_clicked.emit(self.file_path))


# ----------------- BROWSER & EXTENSION HELPER -----------------
def launch_browser_with_extension(url: str, ext_paths: list):
    chrome_candidates = [
        r'C:\Program Files\Google\Chrome\Application\chrome.exe',
        r'C:\Program Files (x86)\Google\Chrome\Application\chrome.exe',
        os.path.expandvars(r'%LOCALAPPDATA%\Google\Chrome\Application\chrome.exe')
    ]
    chrome_exe = None
    for c in chrome_candidates:
        if os.path.exists(c):
            chrome_exe = c
            break

    valid_exts = [os.path.abspath(p) for p in ext_paths if os.path.exists(p)]

    if chrome_exe and valid_exts:
        try:
            ext_arg = ",".join(valid_exts)
            subprocess.Popen([chrome_exe, f"--load-extension={ext_arg}", url])
            return True
        except Exception:
            pass

    webbrowser.open(url)
    return False


# ----------------- ZALO IMPORT DIALOG -----------------
class ZaloImportDialog(QDialog):
    packs_imported = Signal(list)

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Lấy ảnh từ Zalo Web")
        self.setFixedSize(460, 220)
        self.setStyleSheet("""
            QDialog { background-color: #f8fafc; }
            QLabel, QPushButton {
                font-family: 'Segoe UI', -apple-system, BlinkMacSystemFont, Roboto, sans-serif;
            }
        """)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 18, 20, 18)
        layout.setSpacing(12)

        lbl_title = QLabel("Lấy ảnh Zalo chia theo Pack")
        lbl_title.setStyleSheet("font-size: 14px; font-weight: 700; color: #0f172a;")
        layout.addWidget(lbl_title)

        lbl_desc = QLabel("Mở Zalo Web kèm tiện ích gom Pack hoặc nạp file ZIP đã tải về:")
        lbl_desc.setStyleSheet("font-size: 11.5px; color: #64748b;")
        layout.addWidget(lbl_desc)

        btn_open_zalo = QPushButton("Mở Zalo Web (Kèm Extension gom Pack)")
        btn_open_zalo.setCursor(Qt.PointingHandCursor)
        btn_open_zalo.setFixedHeight(36)
        btn_open_zalo.setStyleSheet("""
            QPushButton {
                background-color: #0068ff; color: #ffffff;
                font-size: 12.5px; font-weight: 700;
                border: none; border-radius: 6px;
            }
            QPushButton:hover { background-color: #0052cc; }
        """)
        btn_open_zalo.clicked.connect(self.open_zalo_web)
        layout.addWidget(btn_open_zalo)

        btn_load_zip = QPushButton("Nạp file ZIP Pack ảnh từ máy...")
        btn_load_zip.setCursor(Qt.PointingHandCursor)
        btn_load_zip.setFixedHeight(36)
        btn_load_zip.setStyleSheet("""
            QPushButton {
                background-color: #ffffff; color: #334155;
                font-size: 12px; font-weight: 600;
                border: 1px solid #cbd5e1; border-radius: 6px;
            }
            QPushButton:hover { background-color: #f1f5f9; border-color: #94a3b8; }
        """)
        btn_load_zip.clicked.connect(self.load_zalo_zip)
        layout.addWidget(btn_load_zip)

        layout.addStretch()

        footer = QHBoxLayout()
        footer.addStretch()
        btn_close = QPushButton("Đóng")
        btn_close.setFixedHeight(28)
        btn_close.setFixedWidth(65)
        btn_close.setStyleSheet("""
            QPushButton {
                background: #f1f5f9; color: #475569;
                border: 1px solid #cbd5e1; border-radius: 5px; font-size: 11.5px;
            }
            QPushButton:hover { background: #e2e8f0; }
        """)
        btn_close.clicked.connect(self.reject)
        footer.addWidget(btn_close)
        layout.addLayout(footer)

    def open_zalo_web(self):
        ext_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "extensions", "zalo_pack_collector")
        launch_browser_with_extension("https://chat.zalo.me", [ext_path])
        QMessageBox.information(
            self,
            "Đã mở Zalo Web",
            "Trình duyệt đã mở Zalo Web kèm tiện ích Zalo Pack Collector.\n\n"
            "Hãy bấm biểu tượng tiện ích trên thanh trình duyệt để quét ảnh theo nhân viên & ngày, "
            "sau đó tải file ZIP về và nạp vào phần mềm."
        )

    def load_zalo_zip(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self, "Chọn file ZIP ảnh Zalo", "", "ZIP Files (*.zip *.rar)"
        )
        if not file_path:
            return

        try:
            extract_dir = os.path.join(tempfile.gettempdir(), f"zalo_packs_{datetime.now().strftime('%Y%m%d_%H%M%S')}")
            os.makedirs(extract_dir, exist_ok=True)
            img_paths = []
            with zipfile.ZipFile(file_path, 'r') as zf:
                zf.extractall(extract_dir)
                for root, _, files in os.walk(extract_dir):
                    for fn in files:
                        if fn.lower().endswith(('.jpg', '.jpeg', '.png', '.webp', '.bmp')):
                            img_paths.append(os.path.join(root, fn))

            if img_paths:
                self.packs_imported.emit(img_paths)
                self.accept()
            else:
                QMessageBox.warning(self, "Không tìm thấy ảnh", "File ZIP không chứa hình ảnh hợp lệ!")
        except Exception as e:
            QMessageBox.critical(self, "Lỗi nạp ZIP", f"Không thể đọc file ZIP: {str(e)}")


# ----------------- DROPZONE WIDGET -----------------
class DropZoneWidget(QFrame):
    files_dropped = Signal(list)
    browse_clicked = Signal()

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setAcceptDrops(True)
        self.setFixedHeight(36)
        self.setObjectName("DropZone")
        self.setStyleSheet("""
            QFrame#DropZone {
                background-color: #f8fafc;
                border: 1.5px dashed #cbd5e1;
                border-radius: 6px;
            }
            QFrame#DropZone:hover {
                background-color: #eff6ff;
                border-color: #3b82f6;
            }
        """)

        layout = QHBoxLayout(self)
        layout.setContentsMargins(10, 2, 6, 2)
        layout.setSpacing(8)

        self.prompt_label = QLabel("Kéo thả ảnh vào đây")
        self.prompt_label.setStyleSheet("color: #64748b; font-size: 11.5px; font-weight: 600; border: none; background: transparent;")
        layout.addWidget(self.prompt_label, stretch=1)

        self.btn_browse = QPushButton("Chọn ảnh...")
        self.btn_browse.setCursor(Qt.PointingHandCursor)
        self.btn_browse.setStyleSheet("""
            QPushButton {
                background-color: #ffffff;
                color: #2563eb;
                border: 1px solid #bfdbfe;
                border-radius: 4px;
                padding: 3px 10px;
                font-size: 11.5px;
                font-weight: 600;
            }
            QPushButton:hover {
                background-color: #eff6ff;
                border-color: #2563eb;
            }
        """)
        self.btn_browse.clicked.connect(self.browse_clicked.emit)
        layout.addWidget(self.btn_browse)

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()
            self.setStyleSheet("""
                QFrame#DropZone {
                    background-color: #eff6ff;
                    border: 2px dashed #2563eb;
                    border-radius: 10px;
                }
            """)

    def dragLeaveEvent(self, event):
        self.setStyleSheet("""
            QFrame#DropZone {
                background-color: #f8fafc;
                border: 2px dashed #cbd5e1;
                border-radius: 10px;
            }
        """)

    def dropEvent(self, event):
        self.setStyleSheet("""
            QFrame#DropZone {
                background-color: #f8fafc;
                border: 2px dashed #cbd5e1;
                border-radius: 10px;
            }
        """)
        files = []
        for url in event.mimeData().urls():
            fpath = url.toLocalFile()
            if fpath.lower().endswith(('.jpg', '.jpeg', '.png', '.webp', '.bmp')):
                files.append(fpath)
        if files:
            self.files_dropped.emit(files)


# ----------------- IMAGE VIEWER MODAL (LIGHTBOX) -----------------
class ImageViewerDialog(QDialog):
    def __init__(self, image_path: str, annotated_img: np.ndarray = None, info: dict = None, parent=None):
        super().__init__(parent)
        self.image_path = image_path
        self.info = info or {}
        self.annotated_img = annotated_img
        self.show_tracking = True  # Mặc định bật chế độ tracking màu xanh lá

        # Nếu chưa có ảnh annotated, thử lấy từ cache engine
        if self.annotated_img is None:
            try:
                engine = get_ocr_engine()
                self.annotated_img = engine.get_annotated_preview(image_path)
                if self.annotated_img is None and os.path.exists(image_path):
                    raw_bgr = safe_read_cv2(image_path)
                    if raw_bgr is not None:
                        res = engine.extract_timemark_info(raw_bgr)
                        self.annotated_img = res.get("annotated_image")
                        if not self.info:
                            self.info = res
            except Exception:
                pass

        fn_base = os.path.basename(image_path)
        short_title_fn = fn_base if len(fn_base) <= 30 else (fn_base[:14] + "..." + fn_base[-12:])
        self.setWindowTitle(f"Xem ảnh & Track OCR: {short_title_fn}")
        self.resize(980, 740)
        self.setMinimumSize(700, 500)
        self.setStyleSheet("""
            QDialog {
                background-color: #1e293b;
            }
        """)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(18, 14, 18, 16)
        layout.setSpacing(10)

        # Top Bar
        top_bar = QHBoxLayout()
        top_bar.setSpacing(12)

        lbl_fname = QLabel(short_title_fn)
        lbl_fname.setToolTip(fn_base)
        lbl_fname.setStyleSheet("color: #f8fafc; font-size: 14px; font-weight: 700; border: none; background: transparent;")
        top_bar.addWidget(lbl_fname)

        top_bar.addStretch()

        # Nút chuyển đổi chế độ xem: ROI Tracking Xanh Lá vs Ảnh Gốc
        self.btn_toggle_mode = QPushButton("Đang xem: Vùng ROI & Track chữ")
        self.btn_toggle_mode.setCursor(Qt.PointingHandCursor)
        self.btn_toggle_mode.setStyleSheet("""
            QPushButton {
                background-color: #064e3b;
                color: #a7f3d0;
                border: 1px solid #059669;
                border-radius: 6px;
                padding: 6px 14px;
                font-weight: 700;
                font-size: 12px;
            }
            QPushButton:hover {
                background-color: #047857;
                color: #ffffff;
            }
        """)
        self.btn_toggle_mode.clicked.connect(self.toggle_tracking_mode)
        top_bar.addWidget(self.btn_toggle_mode)

        btn_close = QPushButton("Đóng (Esc)")
        btn_close.setCursor(Qt.PointingHandCursor)
        btn_close.setStyleSheet("""
            QPushButton {
                background-color: #1e293b;
                color: #f8fafc;
                border: 1px solid #334155;
                border-radius: 6px;
                padding: 6px 16px;
                font-weight: 600;
                font-size: 12px;
            }
            QPushButton:hover {
                background-color: #334155;
                border-color: #475569;
            }
        """)
        btn_close.clicked.connect(self.accept)
        top_bar.addWidget(btn_close)
        layout.addLayout(top_bar)

        # Image Scroll Display
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setStyleSheet("""
            QScrollArea {
                border: 1px solid #475569;
                border-radius: 8px;
                background-color: #334155;
            }
        """)

        container = QWidget()
        container.setStyleSheet("background-color: transparent;")
        c_layout = QVBoxLayout(container)
        c_layout.setContentsMargins(10, 10, 10, 10)
        c_layout.setAlignment(Qt.AlignCenter)

        self.img_label = QLabel()
        self.img_label.setAlignment(Qt.AlignCenter)
        self.img_label.setStyleSheet("border: none; background: transparent;")
        c_layout.addWidget(self.img_label)
        scroll.setWidget(container)
        layout.addWidget(scroll, stretch=1)

        # Bottom HUD Info Bar
        info_bar = QFrame()
        info_bar.setStyleSheet("""
            QFrame {
                background-color: #1e293b;
                border: 1px solid #334155;
                border-radius: 8px;
            }
        """)
        ib_layout = QHBoxLayout(info_bar)
        ib_layout.setContentsMargins(14, 8, 14, 8)
        ib_layout.setSpacing(16)

        addr_txt = self.info.get("address") or "Chưa xác định"
        time_txt = f"{self.info.get('date', '--')} {self.info.get('time', '--')}"
        gps_txt = self.info.get("gps_text") or "--"

        lbl_i_addr = QLabel(f"Địa chỉ: {addr_txt}")
        lbl_i_addr.setStyleSheet("color: #38bdf8; font-size: 12px; font-weight: 600; border: none; background: transparent;")
        lbl_i_addr.setWordWrap(True)
        ib_layout.addWidget(lbl_i_addr, stretch=1)

        lbl_i_time = QLabel(f"Thoi gian: {time_txt}")
        lbl_i_time.setStyleSheet("color: #e2e8f0; font-size: 12px; font-weight: 500; border: none; background: transparent;")
        ib_layout.addWidget(lbl_i_time)

        lbl_i_gps = QLabel(f"GPS: {gps_txt}")
        lbl_i_gps.setStyleSheet("color: #a7f3d0; font-size: 12px; font-weight: 500; border: none; background: transparent;")
        ib_layout.addWidget(lbl_i_gps)

        layout.addWidget(info_bar)

        self.update_image_display()

    def update_image_display(self):
        if self.show_tracking and self.annotated_img is not None:
            pix = cv2_to_qpixmap(self.annotated_img, 1800, 1300)
            if not pix.isNull():
                self.img_label.setPixmap(pix)
                self.btn_toggle_mode.setText("Đang xem: Vùng ROI & Track chữ")
                self.btn_toggle_mode.setStyleSheet("""
                    QPushButton {
                        background-color: #064e3b;
                        color: #a7f3d0;
                        border: 1px solid #059669;
                        border-radius: 6px;
                        padding: 6px 14px;
                        font-weight: 700;
                        font-size: 12px;
                    }
                    QPushButton:hover {
                        background-color: #047857;
                        color: #ffffff;
                    }
                """)
                return

        pix = safe_load_pixmap(self.image_path, 1800, 1300)
        if not pix.isNull():
            self.img_label.setPixmap(pix)
        else:
            self.img_label.setText("Không thể tải hình ảnh này")
            self.img_label.setStyleSheet("color: #94a3b8; font-size: 14px; border: none;")
        self.btn_toggle_mode.setText("Đang xem: Ảnh gốc")
        self.btn_toggle_mode.setStyleSheet("""
            QPushButton {
                background-color: #1e293b;
                color: #e2e8f0;
                border: 1px solid #475569;
                border-radius: 6px;
                padding: 6px 14px;
                font-weight: 600;
                font-size: 12px;
            }
            QPushButton:hover {
                background-color: #334155;
            }
        """)

    def toggle_tracking_mode(self):
        self.show_tracking = not self.show_tracking
        self.update_image_display()

    def keyPressEvent(self, event):
        if event.key() == Qt.Key_Escape:
            self.accept()
        else:
            super().keyPressEvent(event)


# ----------------- DETAIL MODAL DIALOG -----------------
class StoreDetailDialog(QDialog):
    def __init__(self, data: dict, all_selected_files: List[str] = None, parent=None):
        super().__init__(parent)
        self.data = data  # Lưu lại để dùng trong open_image_viewer
        self.setWindowTitle("Chi tiết đối soát Timemark")
        self.resize(740, 620)
        self.setStyleSheet("QDialog { background-color: #ffffff; }")

        layout = QVBoxLayout(self)
        layout.setContentsMargins(24, 20, 24, 20)
        layout.setSpacing(14)

        # 1. Header (Tiêu đề + Status Pill)
        top_h = QHBoxLayout()
        lbl_title = QLabel("Thông tin đối soát chi tiết")
        lbl_title.setStyleSheet("font-size: 16px; font-weight: 800; color: #0f172a; border: none;")
        top_h.addWidget(lbl_title)
        top_h.addStretch()

        status = data.get("status", "valid")
        cnt = data.get("count", 1)
        status_pill = QLabel()
        if status == "valid":
            status_pill.setText("Hợp lệ")
            status_pill.setStyleSheet("background-color: #ecfdf5; color: #047857; border: 1px solid #a7f3d0; border-radius: 12px; padding: 4px 14px; font-weight: 700; font-size: 12px;")
        elif status == "duplicate":
            dup_cnt = max(1, cnt - 1)
            status_pill.setText(f"Hợp lệ (Trùng {dup_cnt} ảnh)")
            status_pill.setStyleSheet("background-color: #fffbeb; color: #b45309; border: 1px solid #fde68a; border-radius: 12px; padding: 4px 14px; font-weight: 700; font-size: 12px;")
        else:
            status_pill.setText("Không liên quan")
            status_pill.setStyleSheet("background-color: #fef2f2; color: #b91c1c; border: 1px solid #fecaca; border-radius: 12px; padding: 4px 14px; font-weight: 700; font-size: 12px;")
        top_h.addWidget(status_pill)
        layout.addLayout(top_h)

        # 2. Cảnh báo trùng lặp (nếu có)
        if status == "duplicate":
            dup_cnt = max(1, cnt - 1)
            warn_card = QFrame()
            warn_card.setStyleSheet("background-color: #fffbeb; border: 1px solid #fde68a; border-radius: 8px;")
            wc_l = QHBoxLayout(warn_card)
            wc_l.setContentsMargins(12, 8, 12, 8)
            wc_l.setSpacing(8)
            w_txt = QLabel(f"Cảnh báo trùng lặp: Điểm bán này có {cnt} ảnh ({dup_cnt} ảnh chụp trùng). Hệ thống tính 1 ảnh hợp lệ.")
            w_txt.setStyleSheet("color: #92400e; font-weight: 600; font-size: 12.5px; border: none; background: transparent;")
            wc_l.addWidget(w_txt)
            wc_l.addStretch()
            layout.addWidget(warn_card)

        # 3. Thẻ thông tin chính (Gọn gàng, chữ rõ nét không bị chìm)
        info_card = QFrame()
        info_card.setStyleSheet("background-color: #f8fafc; border: 1px solid #e2e8f0; border-radius: 10px;")
        ic_layout = QVBoxLayout(info_card)
        ic_layout.setContentsMargins(16, 14, 16, 14)
        ic_layout.setSpacing(10)

        # Hàng 1: Nhân viên & Công ty (Chia 2 cột)
        row1 = QHBoxLayout()
        row1.setSpacing(24)

        # Cột Nhân viên SR
        col_sr = QVBoxLayout()
        col_sr.setSpacing(2)
        lbl_sr_head = QLabel("NHÂN VIÊN (SR)")
        lbl_sr_head.setStyleSheet("font-size: 11px; font-weight: 700; color: #64748b; border: none; background: transparent;")
        col_sr.addWidget(lbl_sr_head)
        lbl_sr_val = QLabel(data.get("sr", "--"))
        lbl_sr_val.setStyleSheet("font-size: 12.5px; font-weight: 600; color: #1e293b; border: none; background: transparent;")
        col_sr.addWidget(lbl_sr_val)
        row1.addLayout(col_sr, stretch=1)

        # Cột Công ty
        col_cp = QVBoxLayout()
        col_cp.setSpacing(2)
        lbl_cp_head = QLabel("CÔNG TY")
        lbl_cp_head.setStyleSheet("font-size: 11px; font-weight: 700; color: #64748b; border: none; background: transparent;")
        col_cp.addWidget(lbl_cp_head)
        lbl_cp_val = QLabel(data.get("company", "--"))
        lbl_cp_val.setStyleSheet("font-size: 12.5px; font-weight: 600; color: #1e293b; border: none; background: transparent;")
        col_cp.addWidget(lbl_cp_val)
        row1.addLayout(col_cp, stretch=1)
        ic_layout.addLayout(row1)

        # Đường kẻ phân cách nhẹ 1
        line1 = QFrame()
        line1.setFrameShape(QFrame.HLine)
        line1.setStyleSheet("border: none; border-top: 1px solid #e2e8f0;")
        ic_layout.addWidget(line1)

        # Hàng 2: Địa chỉ điểm bán
        addr_box = QVBoxLayout()
        addr_box.setSpacing(3)
        lbl_addr_head = QLabel("ĐỊA CHỈ ĐIỂM BÁN")
        lbl_addr_head.setStyleSheet("font-size: 11px; font-weight: 700; color: #64748b; letter-spacing: 0.5px; border: none; background: transparent;")
        addr_box.addWidget(lbl_addr_head)

        clean_addr = format_clean_address(data.get("address", ""))
        lbl_addr = QLabel(clean_addr)
        lbl_addr.setWordWrap(True)
        lbl_addr.setStyleSheet("font-size: 13px; font-weight: 700; color: #0f172a; border: none; background: transparent;")
        addr_box.addWidget(lbl_addr)
        ic_layout.addLayout(addr_box)

        # Đường kẻ phân cách nhẹ 2
        line2 = QFrame()
        line2.setFrameShape(QFrame.HLine)
        line2.setStyleSheet("border: none; border-top: 1px solid #e2e8f0;")
        ic_layout.addWidget(line2)

        # Hàng 3: Thời gian & Mã Timemark (Chia 2 cột)
        row3 = QHBoxLayout()
        row3.setSpacing(24)

        # Cột Thời gian
        col_time = QVBoxLayout()
        col_time.setSpacing(2)
        lbl_t_head = QLabel("THỜI GIAN CHỤP")
        lbl_t_head.setStyleSheet("font-size: 11px; font-weight: 700; color: #64748b; border: none; background: transparent;")
        col_time.addWidget(lbl_t_head)
        lbl_t_val = QLabel(f"{data.get('date', '--')}  lúc  {data.get('time', '--')}")
        lbl_t_val.setStyleSheet("font-size: 12.5px; font-weight: 600; color: #1e293b; border: none; background: transparent;")
        col_time.addWidget(lbl_t_val)
        row3.addLayout(col_time, stretch=1)

        # Cột Mã Timemark
        col_tm = QVBoxLayout()
        col_tm.setSpacing(2)
        lbl_tm_head = QLabel("MÃ TIMEMARK")
        lbl_tm_head.setStyleSheet("font-size: 11px; font-weight: 700; color: #64748b; border: none; background: transparent;")
        col_tm.addWidget(lbl_tm_head)
        lbl_tm_val = QLabel(data.get("timemark_code", "--"))
        lbl_tm_val.setStyleSheet("""
            background-color: #eff6ff;
            color: #1d4ed8;
            font-size: 12px;
            font-weight: 700;
            padding: 2px 8px;
            border-radius: 6px;
            border: 1px solid #bfdbfe;
        """)
        col_tm.addWidget(lbl_tm_val, alignment=Qt.AlignLeft)
        row3.addLayout(col_tm, stretch=1)
        ic_layout.addLayout(row3)

        layout.addWidget(info_card)

        # 4. Danh sách hình ảnh (Có thể bấm vào để xem ảnh lớn)
        file_list = data.get("file_list", [])
        top_img_h = QHBoxLayout()
        lbl_sub = QLabel(f"Danh sách ảnh ({len(file_list)} file) — Click vào ảnh để phóng to:")
        lbl_sub.setStyleSheet("font-size: 13px; font-weight: 700; color: #0f172a; border: none;")
        top_img_h.addWidget(lbl_sub)
        top_img_h.addStretch()
        layout.addLayout(top_img_h)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setMinimumHeight(170)
        scroll.setStyleSheet("""
            QScrollArea {
                border: 1px solid #e2e8f0;
                border-radius: 8px;
                background-color: #ffffff;
            }
        """)

        container = QWidget()
        container.setStyleSheet("background-color: #ffffff;")
        c_layout = QVBoxLayout(container)
        c_layout.setContentsMargins(10, 10, 10, 10)
        c_layout.setSpacing(8)

        all_paths = all_selected_files or []

        for idx, fn in enumerate(file_list, 1):
            row_card = QFrame()
            row_card.setStyleSheet("""
                QFrame {
                    background-color: #f8fafc;
                    border: 1px solid #e2e8f0;
                    border-radius: 8px;
                }
                QFrame:hover {
                    background-color: #f1f5f9;
                    border-color: #cbd5e1;
                }
            """)
            r_l = QHBoxLayout(row_card)
            r_l.setContentsMargins(10, 8, 12, 8)
            r_l.setSpacing(14)

            # Tìm đường dẫn file ảnh thật
            match_p = None
            for p in all_paths:
                if os.path.basename(p) == fn:
                    match_p = p
                    break

            # Thumbnail có thể bấm vào để xem lớn
            thumb_box = QLabel()
            thumb_box.setFixedSize(72, 52)
            thumb_box.setStyleSheet("""
                QLabel {
                    border-radius: 6px;
                    border: 1px solid #cbd5e1;
                    background-color: #ffffff;
                }
                QLabel:hover {
                    border: 2px solid #2563eb;
                }
            """)
            thumb_box.setAlignment(Qt.AlignCenter)
            thumb_box.setCursor(Qt.PointingHandCursor)
            thumb_box.setToolTip("Click vào ảnh để phóng to xem chi tiết")

            if match_p:
                pix = safe_load_pixmap(match_p, 70, 50)
                if not pix.isNull():
                    thumb_box.setPixmap(pix)
                # Gán sự kiện click thumbnail mở ảnh lớn
                thumb_box.mousePressEvent = lambda ev, p=match_p: self.open_image_viewer(p)
            r_l.addWidget(thumb_box)

            # Thông tin tên file & tag
            info_col = QVBoxLayout()
            info_col.setSpacing(3)
            info_col.setAlignment(Qt.AlignVCenter)

            name_lbl = QLabel(f"{idx}. {fn}")
            name_lbl.setStyleSheet("color: #0f172a; font-size: 13px; font-weight: 700; border: none; background: transparent;")
            info_col.addWidget(name_lbl)

            if status == "duplicate":
                tag_lbl = QLabel("Ảnh chính (Tính 1)" if idx == 1 else "Ảnh trùng lặp")
                tag_style = """
                    background-color: #ecfdf5; color: #047857; font-size: 11px; font-weight: 700;
                    padding: 2px 8px; border-radius: 6px; border: 1px solid #a7f3d0;
                """ if idx == 1 else """
                    background-color: #fffbeb; color: #b45309; font-size: 11px; font-weight: 700;
                    padding: 2px 8px; border-radius: 6px; border: 1px solid #fde68a;
                """
                tag_lbl.setStyleSheet(tag_style)
                info_col.addWidget(tag_lbl)
            else:
                tag_lbl = QLabel("Ảnh hợp lệ")
                tag_lbl.setStyleSheet("background-color: #ecfdf5; color: #047857; font-size: 11px; font-weight: 700; padding: 2px 8px; border-radius: 6px; border: 1px solid #a7f3d0;")
                info_col.addWidget(tag_lbl)

            r_l.addLayout(info_col, stretch=1)

            # Nút Xem ảnh lớn
            if match_p:
                btn_zoom = QPushButton("Xem ảnh")
                btn_zoom.setCursor(Qt.PointingHandCursor)
                btn_zoom.setFixedSize(80, 30)
                btn_zoom.setStyleSheet("""
                    QPushButton {
                        background-color: #ffffff;
                        color: #2563eb;
                        border: 1px solid #bfdbfe;
                        border-radius: 6px;
                        font-size: 11.5px;
                        font-weight: 600;
                    }
                    QPushButton:hover {
                        background-color: #2563eb;
                        color: #ffffff;
                        border-color: #2563eb;
                    }
                """)
                btn_zoom.clicked.connect(lambda _, p=match_p: self.open_image_viewer(p))
                r_l.addWidget(btn_zoom)

            c_layout.addWidget(row_card)

        c_layout.addStretch()
        scroll.setWidget(container)
        layout.addWidget(scroll, stretch=1)

        # 5. Footer: Nút Đóng
        btn_close = QPushButton("Đóng")
        btn_close.setCursor(Qt.PointingHandCursor)
        btn_close.setFixedSize(80, 32)
        btn_close.setStyleSheet("""
            QPushButton {
                background-color: #f1f5f9;
                color: #334155;
                border: 1px solid #cbd5e1;
                border-radius: 6px;
                font-size: 13px;
                font-weight: 600;
            }
            QPushButton:hover {
                background-color: #e2e8f0;
                color: #0f172a;
            }
        """)
        btn_close.clicked.connect(self.accept)
        layout.addWidget(btn_close, alignment=Qt.AlignRight)

    def open_image_viewer(self, file_path: str):
        if not file_path or not os.path.exists(file_path):
            QMessageBox.information(self, "Thông báo", "Không tìm thấy file ảnh gốc trên máy tính!")
            return
        engine = get_ocr_engine()
        ann = engine.get_annotated_preview(file_path)
        dlg = ImageViewerDialog(file_path, annotated_img=ann, info=self.data, parent=self)
        dlg.exec()


# ----------------- GEMINI ASSISTANT DIALOG -----------------
class GeminiAssistantDialog(QDialog):
    def __init__(self, raw_records: list, on_apply=None, initial_tab: int = 0, parent=None):
        super().__init__(parent)
        self.raw_records = raw_records or []
        self.on_apply = on_apply
        self.setWindowTitle("Trợ lý AI Gemini - Làm sạch dữ liệu & Chấm điểm bán")
        self.resize(820, 640)
        self.setMinimumSize(720, 520)
        self.setStyleSheet("""
            QDialog {
                background-color: #f8fafc;
            }
            QLabel, QPushButton, QTextEdit {
                font-family: 'Segoe UI', -apple-system, BlinkMacSystemFont, Roboto, sans-serif;
            }
        """)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(24, 20, 24, 20)
        layout.setSpacing(14)

        # Header
        hdr = QHBoxLayout()
        hdr_v = QVBoxLayout()
        hdr_v.setSpacing(2)
        hdr_t = QLabel("Trợ lý AI Gemini: Làm sạch JSON & Chấm điểm")
        hdr_t.setStyleSheet("font-size: 15px; font-weight: 800; color: #0f172a;")
        hdr_sub = QLabel("Xuất JSON tên SR, Địa chỉ, Tọa độ, Thời gian để Gemini chuẩn hóa, sau đó nạp lại để chấm điểm chuẩn xác.")
        hdr_sub.setStyleSheet("font-size: 11.5px; color: #64748b;")
        hdr_v.addWidget(hdr_t)
        hdr_v.addWidget(hdr_sub)
        hdr.addLayout(hdr_v)
        hdr.addStretch()
        layout.addLayout(hdr)

        # Tab Widget
        self.tabs = QTabWidget()
        self.tabs.setStyleSheet("""
            QTabWidget::pane {
                border: 1px solid #e2e8f0;
                background-color: #ffffff;
                border-radius: 10px;
                top: -1px;
            }
            QTabBar::tab {
                background: #f1f5f9;
                color: #475569;
                font-weight: 700;
                font-size: 13px;
                padding: 9px 18px;
                margin-right: 4px;
                border-top-left-radius: 8px;
                border-top-right-radius: 8px;
                border: 1px solid #cbd5e1;
                border-bottom: none;
            }
            QTabBar::tab:selected {
                background: #ffffff;
                color: #4338ca;
                border: 1px solid #e2e8f0;
                border-bottom: 2px solid #4f46e5;
            }
            QTabBar::tab:hover:!selected {
                background: #e2e8f0;
            }
        """)

        # --- TAB 1: Xuất JSON & Lấy Prompt ---
        tab_export = QWidget()
        te_l = QVBoxLayout(tab_export)
        te_l.setContentsMargins(18, 16, 18, 16)
        te_l.setSpacing(10)

        te_info = QLabel(f"Đã trích xuất <b>{len(self.raw_records)} ảnh</b> từ watermark Timemark. Bấm <b>Sao chép Prompt + JSON</b> rồi dán thẳng vào Gemini:")
        te_info.setStyleSheet("font-size: 12.5px; color: #334155;")
        te_l.addWidget(te_info)

        self.txt_export = QTextEdit()
        self.txt_export.setReadOnly(True)
        self.txt_export.setStyleSheet("""
            QTextEdit {
                background-color: #f8fafc;
                border: 1px solid #cbd5e1;
                border-radius: 8px;
                font-family: 'Consolas', 'Courier New', monospace;
                font-size: 12px;
                color: #0f172a;
                padding: 8px;
            }
        """)
        json_export_str = json.dumps(self.raw_records, ensure_ascii=False, indent=2)
        self.txt_export.setPlainText(json_export_str)
        te_l.addWidget(self.txt_export)

        # Thanh nút thao tác Tab 1
        te_btns = QHBoxLayout()
        te_btns.setSpacing(10)

        self.btn_copy_prompt = QPushButton("Sao chép Prompt + JSON cho Gemini")
        self.btn_copy_prompt.setCursor(Qt.PointingHandCursor)
        self.btn_copy_prompt.setFixedHeight(38)
        self.btn_copy_prompt.setStyleSheet("""
            QPushButton {
                background-color: #4f46e5;
                color: #ffffff;
                font-size: 13px;
                font-weight: 700;
                border: none;
                border-radius: 7px;
                padding: 0 16px;
            }
            QPushButton:hover { background-color: #4338ca; }
        """)
        self.btn_copy_prompt.clicked.connect(self.copy_prompt_to_clipboard)
        te_btns.addWidget(self.btn_copy_prompt)

        self.btn_save_json = QPushButton("Lưu file JSON ra máy...")
        self.btn_save_json.setCursor(Qt.PointingHandCursor)
        self.btn_save_json.setFixedHeight(38)
        self.btn_save_json.setStyleSheet("""
            QPushButton {
                background-color: #f1f5f9;
                color: #334155;
                border: 1px solid #cbd5e1;
                border-radius: 7px;
                font-size: 13px;
                font-weight: 600;
                padding: 0 16px;
            }
            QPushButton:hover { background-color: #e2e8f0; color: #0f172a; }
        """)
        self.btn_save_json.clicked.connect(self.save_json_file)
        te_btns.addWidget(self.btn_save_json)

        te_btns.addStretch()

        self.lbl_copy_notice = QLabel("")
        self.lbl_copy_notice.setStyleSheet("color: #16a34a; font-weight: 700; font-size: 12px;")
        te_btns.addWidget(self.lbl_copy_notice)

        te_l.addLayout(te_btns)
        self.tabs.addTab(tab_export, "1. Xuất JSON & Prompt")

        # --- TAB 2: Nạp JSON đã làm sạch & Chấm điểm ---
        tab_import = QWidget()
        ti_l = QVBoxLayout(tab_import)
        ti_l.setContentsMargins(18, 16, 18, 16)
        ti_l.setSpacing(10)

        ti_info = QLabel("Dán mảng JSON kết quả từ Gemini hoặc mở file JSON đã làm sạch để tiến hành <b>chấm điểm bán</b>:")
        ti_info.setStyleSheet("font-size: 12.5px; color: #334155;")
        ti_l.addWidget(ti_info)

        ti_action_bar = QHBoxLayout()
        ti_action_bar.setSpacing(8)

        btn_open_file = QPushButton("Chọn file JSON đã làm sạch...")
        btn_open_file.setCursor(Qt.PointingHandCursor)
        btn_open_file.setFixedHeight(34)
        btn_open_file.setStyleSheet("""
            QPushButton {
                background-color: #f8fafc; color: #334155;
                border: 1px solid #cbd5e1; border-radius: 6px;
                font-size: 12.5px; font-weight: 600; padding: 0 14px;
            }
            QPushButton:hover { background-color: #e2e8f0; }
        """)
        btn_open_file.clicked.connect(self.load_cleaned_json_file)
        ti_action_bar.addWidget(btn_open_file)

        btn_paste = QPushButton("Dán từ Clipboard")
        btn_paste.setCursor(Qt.PointingHandCursor)
        btn_paste.setFixedHeight(34)
        btn_paste.setStyleSheet("""
            QPushButton {
                background-color: #f8fafc; color: #334155;
                border: 1px solid #cbd5e1; border-radius: 6px;
                font-size: 12.5px; font-weight: 600; padding: 0 14px;
            }
            QPushButton:hover { background-color: #e2e8f0; }
        """)
        btn_paste.clicked.connect(self.paste_from_clipboard)
        ti_action_bar.addWidget(btn_paste)

        btn_clear = QPushButton("Xóa nội dung")
        btn_clear.setCursor(Qt.PointingHandCursor)
        btn_clear.setFixedHeight(34)
        btn_clear.setStyleSheet("""
            QPushButton {
                background: transparent; color: #ef4444;
                border: none; font-size: 12px; font-weight: 600; padding: 0 8px;
            }
            QPushButton:hover { text-decoration: underline; }
        """)
        btn_clear.clicked.connect(lambda: self.txt_import.clear())
        ti_action_bar.addWidget(btn_clear)
        ti_action_bar.addStretch()
        ti_l.addLayout(ti_action_bar)

        self.txt_import = QTextEdit()
        self.txt_import.setPlaceholderText('Dán kết quả JSON từ Gemini vào đây (hỗ trợ cả khối ```json ... ```):\n[\n  {\n    "id": 1,\n    "filename": "...",\n    "sr": "Nguyễn Văn A",\n    "company": "Công ty X",\n    "address": "ĐT943, Định Mỹ, Thoại Sơn, An Giang",\n    "gps": "10.2934, 105.3421",\n    "date": "06/10/2026",\n    "time": "10:41"\n  }\n]')
        self.txt_import.setStyleSheet("""
            QTextEdit {
                background-color: #ffffff;
                border: 1px solid #cbd5e1;
                border-radius: 8px;
                font-family: 'Consolas', 'Courier New', monospace;
                font-size: 12px;
                color: #0f172a;
                padding: 8px;
            }
            QTextEdit:focus {
                border: 2px solid #4f46e5;
            }
        """)
        ti_l.addWidget(self.txt_import)

        # Nút to chấm điểm bán
        self.btn_evaluate = QPushButton("Bắt đầu chấm điểm từ dữ liệu Gemini")
        self.btn_evaluate.setCursor(Qt.PointingHandCursor)
        self.btn_evaluate.setFixedHeight(40)
        self.btn_evaluate.setStyleSheet("""
            QPushButton {
                background-color: #16a34a;
                color: #ffffff;
                font-size: 13.5px;
                font-weight: 800;
                border: none;
                border-radius: 8px;
                padding: 0 20px;
                letter-spacing: 0.3px;
            }
            QPushButton:hover {
                background-color: #15803d;
            }
        """)
        self.btn_evaluate.clicked.connect(self.process_cleaned_json_and_evaluate)
        ti_l.addWidget(self.btn_evaluate)

        self.tabs.addTab(tab_import, "2. Nạp JSON & Chấm điểm bán")
        layout.addWidget(self.tabs)

        if initial_tab in (0, 1):
            self.tabs.setCurrentIndex(initial_tab)

        # Footer close
        footer = QHBoxLayout()
        footer.addStretch()
        btn_close = QPushButton("Đóng")
        btn_close.setCursor(Qt.PointingHandCursor)
        btn_close.setFixedSize(85, 34)
        btn_close.setStyleSheet("""
            QPushButton {
                background-color: #f1f5f9; color: #475569;
                border: 1px solid #cbd5e1; border-radius: 6px;
                font-size: 13px; font-weight: 600;
            }
            QPushButton:hover { background-color: #e2e8f0; color: #0f172a; }
        """)
        btn_close.clicked.connect(self.reject)
        footer.addWidget(btn_close)
        layout.addLayout(footer)

    def copy_prompt_to_clipboard(self):
        engine = get_ocr_engine()
        prompt = engine.get_gemini_cleaning_prompt(self.raw_records)
        QApplication.clipboard().setText(prompt)
        self.lbl_copy_notice.setText("Đã sao chép Prompt + JSON! Hãy dán (Ctrl+V) vào Gemini.")

    def save_json_file(self):
        default_name = f"diem_ban_raw_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
        save_path, _ = QFileDialog.getSaveFileName(self, "Lưu file JSON cho Gemini", default_name, "JSON Files (*.json)")
        if not save_path:
            return
        try:
            engine = get_ocr_engine()
            engine.export_to_gemini_json(self.raw_records, save_path, include_prompt_file=True)
            QMessageBox.information(
                self, "Đã xuất JSON",
                f"Đã lưu file JSON thành công tại:\n{save_path}\n\nKèm theo file gợi ý Prompt:\n{os.path.splitext(save_path)[0] + '_prompt_gemini.txt'}"
            )
        except Exception as e:
            QMessageBox.critical(self, "Lỗi", f"Không thể lưu file: {str(e)}")

    def load_cleaned_json_file(self):
        file_path, _ = QFileDialog.getOpenFileName(self, "Chọn file JSON đã làm sạch", "", "JSON Files (*.json *.txt)")
        if not file_path:
            return
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                content = f.read()
            self.txt_import.setPlainText(content)
        except Exception as e:
            QMessageBox.critical(self, "Lỗi", f"Không thể đọc file: {str(e)}")

    def paste_from_clipboard(self):
        clip_text = QApplication.clipboard().text()
        if clip_text:
            self.txt_import.setPlainText(clip_text)

    def process_cleaned_json_and_evaluate(self):
        text = self.txt_import.toPlainText().strip()
        if not text:
            QMessageBox.warning(self, "Cảnh báo", "Vui lòng dán hoặc mở nội dung JSON từ Gemini trước khi bấm chấm điểm!")
            return

        # Làm sạch markdown backticks nếu người dùng copy cả ```json ... ```
        if "```" in text:
            m = re.search(r'```(?:json)?\s*([\s\S]*?)\s*```', text)
            if m:
                text = m.group(1).strip()

        try:
            data = json.loads(text)
        except Exception as e:
            QMessageBox.critical(self, "Lỗi định dạng JSON", f"Nội dung không phải là mã JSON hợp lệ:\n{str(e)}")
            return

        if not isinstance(data, list) or not data:
            QMessageBox.warning(self, "Dữ liệu không hợp lệ", "Dữ liệu JSON phải là một mảng (danh sách) các object [ { ... }, { ... } ]!")
            return

        if self.on_apply:
            self.on_apply(data)
            self.accept()


# ----------------- MAIN WINDOW -----------------
class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Anhiu")
        icon_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "app_icon.ico")
        if os.path.exists(icon_path):
            self.setWindowIcon(QIcon(icon_path))
        self.resize(1200, 870)
        self.setMinimumSize(1000, 720)

        self.selected_files: List[str] = []
        self.all_results: Dict[str, Any] = {}
        self.current_filter_status: str = "all"
        self.worker_thread = None
        self.analyzed_previews: List[Dict[str, Any]] = []
        self.current_preview_idx: int = -1
        self.is_roi_zoom: bool = False
        self.current_preview_cv_img = None
        self.current_preview_info = None
        self.current_preview_fn: str = ""
        self.last_raw_ocr_items: List[Dict[str, Any]] = []

        self.setup_ui()

        # Tải trước mô hình PaddleOCR trong nền để sẵn sàng ngay lập tức khi người dùng bấm chạy
        self.preload_thread = EnginePreloadThread()
        self.preload_thread.start()

    def setup_ui(self):
        outer_scroll = QScrollArea()
        outer_scroll.setWidgetResizable(True)
        outer_scroll.setFrameShape(QFrame.NoFrame)
        outer_scroll.setStyleSheet("QScrollArea { border: none; background-color: #f8fafc; }")
        self.setCentralWidget(outer_scroll)

        central = QWidget()
        central.setObjectName("CentralWidget")
        outer_scroll.setWidget(central)

        main_vbox = QVBoxLayout(central)
        main_vbox.setContentsMargins(18, 14, 18, 16)
        main_vbox.setSpacing(10)

        self.setStyleSheet("""
            QMainWindow {
                background-color: #f8fafc;
            }
            QWidget#CentralWidget {
                background-color: #f8fafc;
            }
            QLabel, QPushButton, QLineEdit {
                font-family: 'Segoe UI', -apple-system, BlinkMacSystemFont, Roboto, sans-serif;
            }
        """)

        # ----------------- 1. HEADER ROW -----------------
        header_layout = QHBoxLayout()
        header_layout.setSpacing(12)

        app_title = QLabel("Anhiu OCR")
        app_title.setStyleSheet("font-size: 16px; font-weight: 800; color: #0f172a;")
        header_layout.addWidget(app_title)

        # Mini KPI summary pills
        self.pill_total = QLabel("Tổng: 0")
        self.pill_total.setStyleSheet("background: #f1f5f9; color: #475569; font-size: 11px; font-weight: 600; padding: 2px 8px; border-radius: 4px; border: 1px solid #cbd5e1;")
        self.pill_valid = QLabel("Hợp lệ: 0")
        self.pill_valid.setStyleSheet("background: #ecfdf5; color: #047857; font-size: 11px; font-weight: 700; padding: 2px 8px; border-radius: 4px; border: 1px solid #a7f3d0;")
        self.pill_warn = QLabel("Trùng/Lỗi: 0")
        self.pill_warn.setStyleSheet("background: #fffbeb; color: #b45309; font-size: 11px; font-weight: 700; padding: 2px 8px; border-radius: 4px; border: 1px solid #fde68a;")

        kpi_pills = QHBoxLayout()
        kpi_pills.setSpacing(6)
        kpi_pills.addWidget(self.pill_total)
        kpi_pills.addWidget(self.pill_valid)
        kpi_pills.addWidget(self.pill_warn)
        header_layout.addLayout(kpi_pills)

        header_layout.addStretch()

        self.lbl_selected_count = QLabel("0 ảnh đã chọn")
        self.lbl_selected_count.setStyleSheet("font-size: 11.5px; font-weight: 600; color: #64748b;")
        header_layout.addWidget(self.lbl_selected_count)

        main_vbox.addLayout(header_layout)

        # ----------------- 2. COMPACT TOOLBAR (Single Clean Row) -----------------
        toolbar_card = QFrame()
        toolbar_card.setStyleSheet("""
            QFrame {
                background-color: #ffffff;
                border: 1px solid #e2e8f0;
                border-radius: 8px;
            }
        """)
        tb_layout = QHBoxLayout(toolbar_card)
        tb_layout.setContentsMargins(8, 6, 8, 6)
        tb_layout.setSpacing(8)

        self.dropzone = DropZoneWidget()
        self.dropzone.files_dropped.connect(self.add_files)
        self.dropzone.browse_clicked.connect(self.browse_files)
        tb_layout.addWidget(self.dropzone, 3)

        self.btn_zalo = QPushButton("Lấy ảnh từ Zalo")
        self.btn_zalo.setCursor(Qt.PointingHandCursor)
        self.btn_zalo.setFixedHeight(34)
        self.btn_zalo.setStyleSheet("""
            QPushButton {
                background-color: #eff6ff;
                color: #0068ff;
                border: 1px solid #bfdbfe;
                border-radius: 6px;
                font-size: 12px;
                font-weight: 700;
                padding: 0 12px;
            }
            QPushButton:hover { background-color: #dbeafe; }
        """)
        self.btn_zalo.clicked.connect(self.open_zalo_dialog)
        tb_layout.addWidget(self.btn_zalo)

        self.btn_start = QPushButton("Bắt đầu nhận diện")
        self.btn_start.setCursor(Qt.PointingHandCursor)
        self.btn_start.setFixedHeight(34)
        self.btn_start.setStyleSheet("""
            QPushButton {
                background-color: #2563eb;
                color: #ffffff;
                border: none;
                border-radius: 6px;
                font-size: 12px;
                font-weight: 700;
                padding: 0 16px;
            }
            QPushButton:hover { background-color: #1d4ed8; }
            QPushButton:disabled { background-color: #94a3b8; }
        """)
        self.btn_start.clicked.connect(self.start_ocr_process)
        tb_layout.addWidget(self.btn_start)

        self.btn_open_gemini = QPushButton("Chấm điểm Gemini")
        self.btn_open_gemini.setCursor(Qt.PointingHandCursor)
        self.btn_open_gemini.setFixedHeight(34)
        self.btn_open_gemini.setStyleSheet("""
            QPushButton {
                background-color: #f5f3ff;
                color: #4f46e5;
                border: 1px solid #c7d2fe;
                border-radius: 6px;
                font-size: 12px;
                font-weight: 700;
                padding: 0 12px;
            }
            QPushButton:hover { background-color: #ede9fe; border-color: #818cf8; }
            QPushButton:disabled { color: #94a3b8; border-color: #e2e8f0; background-color: #f8fafc; }
        """)
        self.btn_open_gemini.clicked.connect(self.open_web_gemini_evaluator)
        self.btn_open_gemini.setEnabled(False)
        tb_layout.addWidget(self.btn_open_gemini)

        self.btn_export = QPushButton("Xuất Excel")
        self.btn_export.setCursor(Qt.PointingHandCursor)
        self.btn_export.setFixedHeight(34)
        self.btn_export.setStyleSheet("""
            QPushButton {
                background-color: #f0fdf4;
                color: #16a34a;
                border: 1px solid #bbf7d0;
                border-radius: 6px;
                font-size: 12px;
                font-weight: 700;
                padding: 0 12px;
            }
            QPushButton:hover { background-color: #dcfce7; }
            QPushButton:disabled { color: #94a3b8; border-color: #e2e8f0; background-color: #f8fafc; }
        """)
        self.btn_export.clicked.connect(self.export_to_excel)
        self.btn_export.setEnabled(False)
        tb_layout.addWidget(self.btn_export)

        tb_layout.addStretch()

        self.btn_clear_all = QPushButton("Xóa tất cả")
        self.btn_clear_all.setCursor(Qt.PointingHandCursor)
        self.btn_clear_all.setFixedHeight(34)
        self.btn_clear_all.setStyleSheet("""
            QPushButton {
                background: transparent;
                color: #ef4444;
                border: none;
                font-size: 11.5px;
                font-weight: 600;
                padding: 0 6px;
            }
            QPushButton:hover { text-decoration: underline; color: #dc2626; }
        """)
        self.btn_clear_all.clicked.connect(self.clear_all_files)
        tb_layout.addWidget(self.btn_clear_all)

        main_vbox.addWidget(toolbar_card)

        # Progress bar mỏng 3px
        self.progress_bar = QProgressBar()
        self.progress_bar.setFixedHeight(3)
        self.progress_bar.setTextVisible(False)
        self.progress_bar.setStyleSheet("""
            QProgressBar { background-color: #e2e8f0; border: none; }
            QProgressBar::chunk { background-color: #2563eb; }
        """)
        self.progress_bar.setVisible(False)
        main_vbox.addWidget(self.progress_bar)

        # Thumbnails strip
        self.thumb_scroll = QScrollArea()
        self.thumb_scroll.setFixedHeight(54)
        self.thumb_scroll.setWidgetResizable(True)
        self.thumb_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        self.thumb_scroll.setVerticalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.thumb_scroll.setStyleSheet("QScrollArea { background-color: transparent; border: none; }")
        self.thumb_container = QWidget()
        self.thumb_container.setStyleSheet("background-color: transparent;")
        self.thumb_layout = QHBoxLayout(self.thumb_container)
        self.thumb_layout.setContentsMargins(0, 0, 0, 0)
        self.thumb_layout.setSpacing(6)
        self.thumb_layout.addStretch()
        self.thumb_scroll.setWidget(self.thumb_container)
        self.thumb_scroll.setVisible(False)
        main_vbox.addWidget(self.thumb_scroll)

        # Dummies for backward safety
        self.status_title = QLabel()
        self.status_desc = QLabel()
        self.status_badge = QLabel()
        self.kpi_total_imgs = QLabel("0")
        self.kpi_valid_stores = QLabel("0")
        self.kpi_warning_count = QLabel("0")

        # Quick Filter Tabs
        filter_layout = QHBoxLayout()
        filter_layout.setSpacing(6)

        self.btn_filter_all = QPushButton("Tất cả (0)")
        self.btn_filter_valid = QPushButton("Hợp lệ (0)")
        self.btn_filter_dup = QPushButton("Trùng lặp (0)")
        self.btn_filter_unrel = QPushButton("Không liên quan (0)")

        self.filter_buttons = [
            (self.btn_filter_all, "all"),
            (self.btn_filter_valid, "valid"),
            (self.btn_filter_dup, "duplicate"),
            (self.btn_filter_unrel, "unrelated")
        ]

        for btn, f_key in self.filter_buttons:
            btn.setCursor(Qt.PointingHandCursor)
            btn.setFixedHeight(28)
            btn.clicked.connect(lambda _, k=f_key: self.set_filter(k))
            filter_layout.addWidget(btn)

        self.current_filter_sr = "all"

        filter_layout.addStretch()

        # Dropdown Lọc Nhân Viên (SR)
        lbl_sr_icon = QLabel("Nhân viên:")
        lbl_sr_icon.setStyleSheet("color: #475569; font-size: 11.5px; font-weight: 600; border: none; background: transparent;")
        filter_layout.addWidget(lbl_sr_icon)

        self.combo_filter_sr = QComboBox()
        self.combo_filter_sr.setFixedHeight(28)
        self.combo_filter_sr.setMinimumWidth(160)
        self.combo_filter_sr.setStyleSheet("""
            QComboBox {
                background-color: #ffffff;
                color: #0f172a;
                font-size: 11.5px;
                font-weight: 600;
                border: 1px solid #cbd5e1;
                border-radius: 5px;
                padding: 1px 8px;
            }
            QComboBox:hover {
                border-color: #2563eb;
            }
            QComboBox::drop-down {
                border: none;
                width: 18px;
            }
            QComboBox QAbstractItemView {
                background-color: #ffffff;
                color: #0f172a;
                selection-background-color: #eff6ff;
                selection-color: #1d4ed8;
                border: 1px solid #cbd5e1;
                border-radius: 5px;
                outline: none;
            }
        """)
        self.combo_filter_sr.addItem("Tất cả nhân viên", "all")
        self.combo_filter_sr.currentIndexChanged.connect(self.on_sr_filter_changed)
        filter_layout.addWidget(self.combo_filter_sr)

        main_vbox.addLayout(filter_layout)
        self.apply_filter_tab_styles()

        # Dải Thẻ Chip Nhân Viên (Bấm trực quan để lọc nhanh theo nhân viên)
        self.staff_chip_container = QWidget()
        self.staff_chip_container.setVisible(False)
        self.staff_chip_layout = QHBoxLayout(self.staff_chip_container)
        self.staff_chip_layout.setContentsMargins(0, 2, 0, 4)
        self.staff_chip_layout.setSpacing(6)
        self.staff_chip_layout.setAlignment(Qt.AlignLeft)
        main_vbox.addWidget(self.staff_chip_container)

        # Search Bar
        self.search_box = QLineEdit()
        self.search_box.setPlaceholderText("Tìm kiếm nhanh theo địa chỉ, nhân viên (SR), công ty, mã timecode...")
        self.search_box.setFixedHeight(32)
        self.search_box.setStyleSheet("""
            QLineEdit {
                background-color: #ffffff;
                border: 1px solid #cbd5e1;
                border-radius: 6px;
                padding: 0 12px;
                font-size: 12px;
                color: #0f172a;
            }
            QLineEdit:focus {
                border: 1.5px solid #2563eb;
                background-color: #ffffff;
            }
        """)
        self.search_box.textChanged.connect(self.filter_table)
        main_vbox.addWidget(self.search_box)

        # ----------------- 4. DATA TABLE (9 COLUMNS) -----------------
        self.table = QTableWidget()
        self.table.setColumnCount(9)
        self.table.setHorizontalHeaderLabels([
            "#", "Thời gian", "Nhân viên (SR)", "Công ty", "Địa chỉ điểm bán", "Mã Timemark", "Hình ảnh", "Trạng thái", "Thao tác"
        ])
        self.table.verticalHeader().setVisible(False)
        self.table.setAlternatingRowColors(True)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.NoSelection)
        self.table.setShowGrid(False)
        self.table.setFocusPolicy(Qt.NoFocus)
        self.table.cellClicked.connect(self.on_table_cell_clicked)
        self.table.cellDoubleClicked.connect(lambda r, c: self.on_table_row_double_clicked(r))

        self.table.verticalHeader().setDefaultSectionSize(40)
        self.table.verticalHeader().setSectionResizeMode(QHeaderView.Fixed)

        header = self.table.horizontalHeader()
        header.setFixedHeight(32)
        header.setSectionResizeMode(0, QHeaderView.Fixed)
        header.resizeSection(0, 40)    # STT
        header.setSectionResizeMode(1, QHeaderView.Fixed)
        header.resizeSection(1, 110)   # Thời gian
        header.setSectionResizeMode(2, QHeaderView.Fixed)
        header.resizeSection(2, 125)   # Nhân viên (SR)
        header.setSectionResizeMode(3, QHeaderView.Fixed)
        header.resizeSection(3, 130)   # Công ty
        header.setSectionResizeMode(4, QHeaderView.Stretch)           # Địa chỉ (Tự co giãn cân đối)
        header.setSectionResizeMode(5, QHeaderView.Fixed)
        header.resizeSection(5, 120)   # Mã Timemark (Timecode)
        header.setSectionResizeMode(6, QHeaderView.Fixed)
        header.resizeSection(6, 100)   # Ảnh
        header.setSectionResizeMode(7, QHeaderView.Fixed)
        header.resizeSection(7, 140)   # Trạng thái
        header.setSectionResizeMode(8, QHeaderView.Fixed)
        header.resizeSection(8, 74)    # Thao tác

        self.table.setStyleSheet("""
            QTableWidget {
                background-color: #ffffff;
                alternate-background-color: #fafbfc;
                border: 1px solid #e2e8f0;
                border-radius: 8px;
                font-size: 12px;
                outline: none;
            }
            QHeaderView::section {
                background-color: #f8fafc;
                color: #475569;
                font-weight: 700;
                font-size: 11px;
                border: none;
                border-bottom: 2px solid #e2e8f0;
                padding: 3px 6px;
            }
            QTableWidget::item {
                border-bottom: 1px solid #f1f5f9;
                color: #1e293b;
            }
            QTableWidget::item:hover {
                background-color: #f1f5f9;
            }
            QScrollBar:vertical {
                border: none;
                background: #f8fafc;
                width: 7px;
                border-radius: 3px;
            }
            QScrollBar::handle:vertical {
                background: #cbd5e1;
                border-radius: 3px;
                min-height: 20px;
            }
            QScrollBar::handle:vertical:hover {
                background: #94a3b8;
            }
        """)
        self.table.setMinimumHeight(240)
        main_vbox.addWidget(self.table)

        self.lbl_footer = QLabel("Chưa có dữ liệu phân tích")
        self.lbl_footer.setStyleSheet("font-size: 12px; color: #94a3b8; font-weight: 500; border: none;")
        main_vbox.addWidget(self.lbl_footer)

        self.update_file_selection_ui()

    # ----------------- UI UPDATES & EVENTS -----------------
    def apply_filter_tab_styles(self):
        active_style = """
            QPushButton {
                background-color: #2563eb;
                color: #ffffff;
                font-weight: 700;
                font-size: 12px;
                border-radius: 6px;
                padding: 0 14px;
                border: none;
            }
        """
        inactive_style = """
            QPushButton {
                background-color: #f1f5f9;
                color: #475569;
                font-weight: 600;
                font-size: 12px;
                border-radius: 6px;
                padding: 0 14px;
                border: 1px solid #e2e8f0;
            }
            QPushButton:hover {
                background-color: #e2e8f0;
            }
        """
        for btn, k in self.filter_buttons:
            if k == self.current_filter_status:
                btn.setStyleSheet(active_style)
            else:
                btn.setStyleSheet(inactive_style)

    def set_filter(self, filter_key: str):
        self.current_filter_status = filter_key
        self.apply_filter_tab_styles()
        self.filter_table(self.search_box.text())

    def add_files(self, file_paths: List[str]):
        added = 0
        for p in file_paths:
            if p not in self.selected_files:
                self.selected_files.append(p)
                added += 1
        if added > 0:
            self.refresh_thumbnails()
            self.update_file_selection_ui()

    def browse_files(self):
        files, _ = QFileDialog.getOpenFileNames(
            self,
            "Chọn ảnh Timemark",
            "",
            "Image Files (*.jpg *.jpeg *.png *.webp *.bmp)"
        )
        if files:
            self.add_files(files)

    def remove_file(self, file_path: str):
        if file_path in self.selected_files:
            self.selected_files.remove(file_path)
            self.refresh_thumbnails()
            self.update_file_selection_ui()

    def clear_all_files(self):
        self.selected_files.clear()
        self.all_results = {}
        self.last_raw_ocr_items.clear()
        self.table.setRowCount(0)
        self.lbl_selected_count.setText("0 ảnh đã chọn")
        self.pill_total.setText("Tổng: 0")
        self.pill_valid.setText("Hợp lệ: 0")
        self.pill_warn.setText("Trùng/Lỗi: 0")
        self.btn_filter_all.setText("Tất cả (0)")
        self.btn_filter_valid.setText("Hợp lệ (0)")
        self.btn_filter_dup.setText("Trùng lặp (0)")
        self.btn_filter_unrel.setText("Không liên quan (0)")
        self.btn_export.setEnabled(False)
        self.btn_open_gemini.setEnabled(False)
        self.progress_bar.setVisible(False)
        self.lbl_footer.setText("Chưa có dữ liệu phân tích")
        self.refresh_thumbnails()
        self.update_file_selection_ui()

    def refresh_thumbnails(self):
        while self.thumb_layout.count() > 0:
            item = self.thumb_layout.takeAt(0)
            widget = item.widget()
            if widget:
                widget.deleteLater()

        for fp in self.selected_files:
            card = ThumbnailCard(fp)
            card.delete_clicked.connect(self.remove_file)
            self.thumb_layout.addWidget(card)

        self.thumb_layout.addStretch()

    def update_file_selection_ui(self):
        n = len(self.selected_files)
        self.lbl_selected_count.setText(f"Đã chọn {n} ảnh")
        has_files = (n > 0)
        self.btn_clear_all.setVisible(has_files)
        self.thumb_scroll.setVisible(has_files)
        self.btn_start.setEnabled(has_files)

    def switch_right_tab(self, tab_name: str):
        pass

    def toggle_roi_zoom(self):
        pass

    def open_current_fullscreen_preview(self):
        pass

    def adjust_preview_frame_aspect(self, cv_img: np.ndarray):
        pass

    def update_live_preview(self, cv_img: np.ndarray, info: dict, fn: str = ""):
        pass

    def show_prev_preview(self):
        pass

    def show_next_preview(self):
        pass

    def open_zalo_dialog(self):
        dlg = ZaloImportDialog(parent=self)
        dlg.packs_imported.connect(self.add_files)
        dlg.exec()

    def on_table_cell_clicked(self, row: int, col: int):
        if not hasattr(self, 'all_results') or not self.all_results:
            return
        details = self.all_results.get("details", [])
        if row < 0 or row >= len(details):
            return
        row_data = details[row]
        file_list = row_data.get("file_list", [])
        if not file_list:
            return

        target_fn = file_list[0]
        match_path = None
        for p in self.selected_files:
            if os.path.basename(p) == target_fn:
                match_path = p
                break

        # Nếu click vào cột Hình ảnh (cột 6), mở ngay ImageViewerDialog phóng to
        if col == 6 and match_path:
            self.open_full_image_viewer(match_path, row_data)
            return

    def on_table_row_double_clicked(self, row: int):
        if not hasattr(self, 'all_results') or not self.all_results:
            return
        details = self.all_results.get("details", [])
        if row < 0 or row >= len(details):
            return
        row_data = details[row]
        file_list = row_data.get("file_list", [])
        if file_list:
            for p in self.selected_files:
                if os.path.basename(p) == file_list[0]:
                    self.open_full_image_viewer(p, row_data)
                    return

    def open_full_image_viewer(self, file_path: str, row_data: dict = None):
        if not file_path or not os.path.exists(file_path):
            QMessageBox.information(self, "Thông báo", "Không tìm thấy file ảnh gốc trên máy tính!")
            return
        engine = get_ocr_engine()
        ann = engine.get_annotated_preview(file_path)
        dlg = ImageViewerDialog(file_path, annotated_img=ann, info=row_data, parent=self)
        dlg.exec()

    def start_ocr_process(self):
        if not self.selected_files:
            QMessageBox.warning(self, "Thông báo", "Vui lòng chọn ít nhất 1 ảnh trước khi bắt đầu!")
            return

        self.btn_start.setEnabled(False)
        self.dropzone.setEnabled(False)
        self.btn_clear_all.setEnabled(False)
        self.btn_export.setEnabled(False)
        self.btn_open_gemini.setEnabled(False)

        self.lbl_selected_count.setText("Đang khởi động AI...")
        self.progress_bar.setVisible(True)
        self.progress_bar.setValue(10)

        self.worker_thread = OCRWorkerThread(self.selected_files)
        self.worker_thread.progress_update.connect(self.on_worker_progress)
        self.worker_thread.finished_success.connect(self.on_worker_success)
        self.worker_thread.finished_error.connect(self.on_worker_error)
        self.worker_thread.start()

    def on_worker_progress(self, data: dict):
        curr = data.get("progress_current", 0)
        tot = data.get("progress_total", 0)
        task_idx = data.get("task_index", 0)

        if task_idx == 2:
            if tot > 0:
                pct = 15 + int((curr / tot) * 70)
                self.progress_bar.setValue(min(pct, 85))
                self.lbl_selected_count.setText(f"Đang nhận diện: {curr}/{tot} ảnh...")
            else:
                self.progress_bar.setValue(40)
        elif task_idx == 3:
            self.progress_bar.setValue(90)
            self.lbl_selected_count.setText("Đang gom nhóm và đối soát điểm bán...")
        elif task_idx == 4:
            self.progress_bar.setValue(100)

    def on_worker_success(self, results: dict):
        self.all_results = results
        self.last_raw_ocr_items = results.get("raw_ocr_results", [])
        self.btn_start.setEnabled(True)
        self.dropzone.setEnabled(True)
        self.btn_clear_all.setEnabled(True)
        self.btn_export.setEnabled(True)
        self.btn_open_gemini.setEnabled(True)
        self.progress_bar.setVisible(False)

        summary = results.get("summary", {})
        total_imgs = summary.get("total_images", 0)
        valid_stores = summary.get("valid_stores", 0)
        warning_cnt = summary.get("warning_total", 0)

        self.pill_total.setText(f"Tổng: {total_imgs}")
        self.pill_valid.setText(f"Hợp lệ: {valid_stores}")
        self.pill_warn.setText(f"Trùng/Lỗi: {warning_cnt}")
        self.lbl_selected_count.setText(f"{total_imgs} ảnh ({valid_stores} điểm bán)")

        details = results.get("details", [])
        c_all = len(details)
        c_valid = sum(1 for d in details if d.get("status") in ("valid", "duplicate"))
        c_dup = sum(1 for d in details if d.get("status") == "duplicate")
        c_unrel = sum(1 for d in details if d.get("status") == "unrelated")

        self.btn_filter_all.setText(f"Tất cả ({c_all})")
        self.btn_filter_valid.setText(f"Hợp lệ ({c_valid})")
        self.btn_filter_dup.setText(f"Trùng lặp ({c_dup})")
        self.btn_filter_unrel.setText(f"Không liên quan ({c_unrel})")

        # Cập nhật danh sách Nhân viên vào Dropdown và Dải Thẻ Chip
        staff_summary = summary.get("staff_summary", {})
        self.combo_filter_sr.blockSignals(True)
        self.combo_filter_sr.clear()
        self.combo_filter_sr.addItem(f"Tất cả nhân viên ({valid_stores} điểm)", "all")

        # Xóa các chip cũ
        while self.staff_chip_layout.count():
            item = self.staff_chip_layout.takeAt(0)
            w = item.widget()
            if w:
                w.deleteLater()

        self.staff_chips = []

        # Chip 'Tất cả'
        btn_chip_all = QPushButton(f"Tất cả ({valid_stores} điểm)")
        btn_chip_all.setCursor(Qt.PointingHandCursor)
        btn_chip_all.clicked.connect(lambda _, k="all": self.set_staff_filter_from_chip(k))
        self.staff_chip_layout.addWidget(btn_chip_all)
        self.staff_chips.append((btn_chip_all, "all"))

        for sr_name, s_data in staff_summary.items():
            st_cnt = s_data.get("total_stores", 0)
            p_cnt = s_data.get("total_photos", 0)
            self.combo_filter_sr.addItem(f"{sr_name} ({st_cnt} điểm - {p_cnt} ảnh)", sr_name)

            btn_chip = QPushButton(f"{sr_name}: {st_cnt} điểm ({p_cnt} ảnh)")
            btn_chip.setCursor(Qt.PointingHandCursor)
            btn_chip.clicked.connect(lambda _, k=sr_name: self.set_staff_filter_from_chip(k))
            self.staff_chip_layout.addWidget(btn_chip)
            self.staff_chips.append((btn_chip, sr_name))

        self.staff_chip_layout.addStretch()
        self.staff_chip_container.setVisible(len(staff_summary) > 0)
        self.combo_filter_sr.blockSignals(False)
        self.current_filter_sr = "all"
        self.update_staff_chips_style()

        self.populate_table(details)

    def on_worker_error(self, err_msg: str):
        self.btn_start.setEnabled(True)
        self.dropzone.setEnabled(True)
        self.btn_clear_all.setEnabled(True)
        self.progress_bar.setVisible(False)
        self.lbl_selected_count.setText("Lỗi nhận diện!")
        QMessageBox.critical(self, "Lỗi phân tích", f"Không thể xử lý ảnh:\n{err_msg}")

    def populate_table(self, rows: List[Dict[str, Any]]):
        self.table.setRowCount(len(rows))

        for r_idx, row_data in enumerate(rows):
            self.table.setRowHeight(r_idx, 62)

            stt = f"{row_data.get('stt', r_idx + 1):02d}"
            date_str = row_data.get("date", "--/--/----")
            time_str = row_data.get("time", "--:--")
            raw_address = row_data.get("address", "")
            address = format_clean_address(raw_address)
            sr = row_data.get("sr", "--")
            company = row_data.get("company", "--")
            tm_code = row_data.get("timemark_code", "--")
            count = row_data.get("count", 1)
            file_list = row_data.get("file_list", [])
            status = row_data.get("status", "valid")

            # Chuỗi tìm kiếm toàn diện cho dòng (bao gồm cả mã timecode)
            search_text = f"{address} {sr} {company} {tm_code} {date_str} {time_str}".lower()

            # 0. STT (Lưu UserRole cho lọc & tìm kiếm để tránh lỗi vẽ đè chữ của QTableWidget)
            stt_item = QTableWidgetItem(stt)
            stt_item.setTextAlignment(Qt.AlignCenter)
            stt_item.setForeground(QColor("#64748b"))
            stt_font = QFont("Segoe UI", 10)
            stt_font.setWeight(QFont.DemiBold)
            stt_item.setFont(stt_font)
            stt_item.setData(Qt.UserRole, status)
            stt_item.setData(Qt.UserRole + 1, search_text)
            stt_item.setData(Qt.UserRole + 2, sr)
            self.table.setItem(r_idx, 0, stt_item)

            # 1. Thời gian (Ngày + Giờ xếp tầng)
            self.table.setItem(r_idx, 1, QTableWidgetItem(""))
            time_widget = QWidget()
            tw_l = QVBoxLayout(time_widget)
            tw_l.setContentsMargins(2, 2, 2, 2)
            tw_l.setSpacing(1)
            tw_l.setAlignment(Qt.AlignCenter)

            d_lbl = QLabel(date_str)
            d_lbl.setStyleSheet("font-size: 11px; font-weight: 700; color: #1e293b; border: none; background: transparent;")
            d_lbl.setAlignment(Qt.AlignCenter)
            tw_l.addWidget(d_lbl)

            t_badge = QLabel(time_str)
            t_badge.setStyleSheet("""
                background-color: #f1f5f9;
                color: #475569;
                font-size: 10px;
                font-weight: 600;
                padding: 1px 6px;
                border-radius: 4px;
                border: 1px solid #e2e8f0;
            """)
            t_badge.setAlignment(Qt.AlignCenter)
            tw_l.addWidget(t_badge)
            self.table.setCellWidget(r_idx, 1, time_widget)

            # 2. Nhân viên SR (Tên nhân viên sạch)
            self.table.setItem(r_idx, 2, QTableWidgetItem(""))
            sr_widget = QWidget()
            sw_l = QHBoxLayout(sr_widget)
            sw_l.setContentsMargins(8, 2, 8, 2)
            sw_l.setAlignment(Qt.AlignVCenter | Qt.AlignLeft)

            if sr and sr != "--":
                name_lbl = QLabel(sr)
                name_lbl.setStyleSheet("font-size: 11.5px; font-weight: 600; color: #334155; border: none; background: transparent;")
                sw_l.addWidget(name_lbl)
            else:
                dash_lbl = QLabel("--")
                dash_lbl.setStyleSheet("color: #94a3b8; font-size: 11.5px; border: none; background: transparent;")
                sw_l.addWidget(dash_lbl)
            sw_l.addStretch()
            self.table.setCellWidget(r_idx, 2, sr_widget)

            # 3. Công ty
            self.table.setItem(r_idx, 3, QTableWidgetItem(""))
            cp_widget = QWidget()
            cw_l = QHBoxLayout(cp_widget)
            cw_l.setContentsMargins(4, 2, 4, 2)
            cw_l.setSpacing(4)
            cw_l.setAlignment(Qt.AlignCenter)

            if company and company != "--":
                cp_badge = QLabel(company)
                cp_badge.setStyleSheet("""
                    background-color: #f8fafc;
                    color: #334155;
                    font-size: 10.5px;
                    font-weight: 600;
                    padding: 2px 8px;
                    border-radius: 5px;
                    border: 1px solid #e2e8f0;
                """)
                cp_badge.setAlignment(Qt.AlignCenter)
                cw_l.addWidget(cp_badge)
            else:
                dash_cp = QLabel("--")
                dash_cp.setStyleSheet("color: #94a3b8; font-size: 11.5px; border: none; background: transparent;")
                cw_l.addWidget(dash_cp)
            self.table.setCellWidget(r_idx, 3, cp_widget)

            # 4. Địa chỉ điểm bán
            self.table.setItem(r_idx, 4, QTableWidgetItem(""))
            addr_widget = QWidget()
            aw_l = QHBoxLayout(addr_widget)
            aw_l.setContentsMargins(8, 2, 8, 2)
            aw_l.setAlignment(Qt.AlignVCenter | Qt.AlignLeft)

            txt_addr = QLabel(address)
            txt_addr.setWordWrap(True)
            if status == "unrelated":
                txt_addr.setStyleSheet("font-size: 11px; color: #b91c1c; font-style: italic; font-weight: 500; border: none; background: transparent;")
            else:
                txt_addr.setStyleSheet("font-size: 11.5px; color: #0f172a; font-weight: 600; border: none; background: transparent;")
            txt_addr.setToolTip(address)
            aw_l.addWidget(txt_addr, stretch=1)
            self.table.setCellWidget(r_idx, 4, addr_widget)

            # 5. Mã Timemark / Timecode
            self.table.setItem(r_idx, 5, QTableWidgetItem(""))
            tm_widget = QWidget()
            tm_l = QHBoxLayout(tm_widget)
            tm_l.setContentsMargins(4, 2, 4, 2)
            tm_l.setAlignment(Qt.AlignCenter)

            if tm_code and tm_code != "--":
                lbl_tm = QLabel(tm_code)
                lbl_tm.setStyleSheet("""
                    background-color: #eff6ff;
                    color: #1d4ed8;
                    font-size: 10.5px;
                    font-weight: 700;
                    padding: 2px 6px;
                    border-radius: 4px;
                    border: 1px solid #bfdbfe;
                """)
                lbl_tm.setAlignment(Qt.AlignCenter)
                tm_l.addWidget(lbl_tm)
            else:
                dash_tm = QLabel("--")
                dash_tm.setStyleSheet("color: #94a3b8; font-size: 11.5px; border: none; background: transparent;")
                dash_tm.setAlignment(Qt.AlignCenter)
                tm_l.addWidget(dash_tm)
            self.table.setCellWidget(r_idx, 5, tm_widget)

            # 6. Hình ảnh
            self.table.setItem(r_idx, 6, QTableWidgetItem(""))
            badge_widget = QWidget()
            bw_l = QHBoxLayout(badge_widget)
            bw_l.setContentsMargins(4, 2, 4, 2)
            bw_l.setSpacing(6)
            bw_l.setAlignment(Qt.AlignCenter)

            first_file = file_list[0] if file_list else ""
            match_path = None
            for p in self.selected_files:
                if os.path.basename(p) == first_file:
                    match_path = p
                    break

            thumb_lbl = QLabel()
            thumb_lbl.setFixedSize(36, 26)
            thumb_lbl.setCursor(Qt.PointingHandCursor)
            thumb_lbl.setStyleSheet("border-radius: 4px; border: 1px solid #cbd5e1; background-color: #f1f5f9;")
            thumb_lbl.setAlignment(Qt.AlignCenter)
            if match_path:
                pix = safe_load_pixmap(match_path, 34, 24)
                if not pix.isNull():
                    thumb_lbl.setPixmap(pix)
                thumb_lbl.mousePressEvent = lambda ev, p=match_path, d=row_data: self.open_full_image_viewer(p, d)
            bw_l.addWidget(thumb_lbl)

            if count > 1:
                cnt_style = """
                    background-color: #eff6ff;
                    color: #2563eb;
                    font-size: 10.5px;
                    font-weight: 700;
                    padding: 2px 6px;
                    border-radius: 8px;
                    border: 1px solid #bfdbfe;
                """
            else:
                cnt_style = """
                    background-color: #f1f5f9;
                    color: #475569;
                    font-size: 10.5px;
                    font-weight: 600;
                    padding: 2px 6px;
                    border-radius: 8px;
                    border: 1px solid #e2e8f0;
                """
            cnt_pill = QLabel(f"{count} ảnh")
            cnt_pill.setStyleSheet(cnt_style)
            bw_l.addWidget(cnt_pill)
            self.table.setCellWidget(r_idx, 6, badge_widget)

            # 7. Trạng thái
            self.table.setItem(r_idx, 7, QTableWidgetItem(""))
            status_container = QWidget()
            sc_l = QHBoxLayout(status_container)
            sc_l.setContentsMargins(2, 2, 2, 2)
            sc_l.setAlignment(Qt.AlignCenter)

            stat_lbl = QLabel()
            if status == "valid":
                stat_lbl.setText("Hợp lệ")
                stat_lbl.setStyleSheet("""
                    background-color: #ecfdf5;
                    color: #047857;
                    border: 1px solid #a7f3d0;
                    font-size: 10.5px;
                    font-weight: 700;
                    padding: 3px 8px;
                    border-radius: 10px;
                """)
            elif status == "duplicate":
                dup_cnt = max(1, count - 1)
                stat_lbl.setText(f"Hợp lệ (Trùng {dup_cnt} ảnh)")
                stat_lbl.setToolTip(f"Điểm bán này được tính 1 ảnh hợp lệ. Cảnh báo phát hiện {dup_cnt} ảnh chụp trùng.")
                stat_lbl.setStyleSheet("""
                    background-color: #fffbeb;
                    color: #b45309;
                    border: 1px solid #fde68a;
                    font-size: 10.5px;
                    font-weight: 700;
                    padding: 3px 8px;
                    border-radius: 10px;
                """)
            else:
                stat_lbl.setText("Không liên quan")
                stat_lbl.setStyleSheet("""
                    background-color: #fef2f2;
                    color: #b91c1c;
                    border: 1px solid #fecaca;
                    font-size: 10.5px;
                    font-weight: 700;
                    padding: 3px 8px;
                    border-radius: 10px;
                """)
            sc_l.addWidget(stat_lbl)
            self.table.setCellWidget(r_idx, 7, status_container)

            # 8. Nút Thao tác
            self.table.setItem(r_idx, 8, QTableWidgetItem(""))
            btn_container = QWidget()
            bc_l = QHBoxLayout(btn_container)
            bc_l.setContentsMargins(2, 2, 2, 2)
            bc_l.setAlignment(Qt.AlignCenter)

            btn_view = QPushButton("Chi tiết")
            btn_view.setCursor(Qt.PointingHandCursor)
            btn_view.setFixedSize(64, 24)
            btn_view.setStyleSheet("""
                QPushButton {
                    background-color: #ffffff;
                    color: #2563eb;
                    border: 1px solid #bfdbfe;
                    border-radius: 5px;
                    font-size: 11px;
                    font-weight: 600;
                }
                QPushButton:hover {
                    background-color: #2563eb;
                    color: #ffffff;
                    border-color: #2563eb;
                }
            """)
            btn_view.clicked.connect(lambda _, d=row_data: self.show_detail(d))
            bc_l.addWidget(btn_view)
            self.table.setCellWidget(r_idx, 8, btn_container)

        self.filter_table(self.search_box.text())

    def show_detail(self, row_data: dict):
        dlg = StoreDetailDialog(row_data, self.selected_files, self)
        dlg.exec()

    def on_sr_filter_changed(self):
        self.current_filter_sr = self.combo_filter_sr.currentData() or "all"
        self.update_staff_chips_style()
        self.filter_table(self.search_box.text())

    def update_staff_chips_style(self):
        if not hasattr(self, 'staff_chips'):
            return
        for btn, sr_val in self.staff_chips:
            is_active = (self.current_filter_sr == sr_val)
            if is_active:
                btn.setStyleSheet("""
                    QPushButton {
                        background-color: #1d4ed8;
                        color: #ffffff;
                        font-size: 11.5px;
                        font-weight: 700;
                        border: 1px solid #1e40af;
                        border-radius: 6px;
                        padding: 3px 10px;
                    }
                """)
            else:
                btn.setStyleSheet("""
                    QPushButton {
                        background-color: #f1f5f9;
                        color: #334155;
                        font-size: 11.5px;
                        font-weight: 600;
                        border: 1px solid #cbd5e1;
                        border-radius: 6px;
                        padding: 3px 10px;
                    }
                    QPushButton:hover {
                        background-color: #e2e8f0;
                        color: #0f172a;
                    }
                """)

    def set_staff_filter_from_chip(self, sr_key: str):
        # Đồng bộ chuyển index của combo_filter_sr
        for i in range(self.combo_filter_sr.count()):
            if self.combo_filter_sr.itemData(i) == sr_key:
                self.combo_filter_sr.setCurrentIndex(i)
                return
        self.current_filter_sr = sr_key
        self.update_staff_chips_style()
        self.filter_table(self.search_box.text())

    def filter_table(self, query: str):
        q = query.strip().lower()
        visible_count = 0
        total = self.table.rowCount()

        for row in range(total):
            item_meta = self.table.item(row, 0)
            if not item_meta:
                self.table.setRowHidden(row, False)
                visible_count += 1
                continue

            row_status = item_meta.data(Qt.UserRole)
            text_all = (item_meta.data(Qt.UserRole + 1) or "").lower()
            row_sr = (item_meta.data(Qt.UserRole + 2) or "").strip()

            # 1. Khớp trạng thái (Tất cả / Hợp lệ / Trùng lặp / Không liên quan)
            if self.current_filter_status == "all":
                status_match = True
            elif self.current_filter_status == "valid":
                status_match = (row_status in ("valid", "duplicate"))
            elif self.current_filter_status == "duplicate":
                status_match = (row_status == "duplicate")
            elif self.current_filter_status == "unrelated":
                status_match = (row_status == "unrelated")
            else:
                status_match = True

            # 2. Khớp từ khóa tìm kiếm
            text_match = (q in text_all) if q else True

            # 3. Khớp Nhân viên (SR)
            if self.current_filter_sr == "all":
                sr_match = True
            else:
                sr_match = (row_sr == self.current_filter_sr)

            matched = status_match and text_match and sr_match
            self.table.setRowHidden(row, not matched)
            if matched:
                visible_count += 1

        self.lbl_footer.setText(f"Hiển thị {visible_count} trong tổng số {total} kết quả")

    def export_to_excel(self):
        details = self.all_results.get("details", [])
        if not details:
            QMessageBox.information(self, "Thông báo", "Chưa có kết quả để xuất file!")
            return

        default_name = f"diem_ban_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        save_path, _ = QFileDialog.getSaveFileName(
            self,
            "Lưu báo cáo Excel",
            default_name,
            "Excel Workbook (*.xlsx)"
        )
        if not save_path:
            return

        try:
            # Điểm bán hợp lệ bao gồm cả điểm trùng ảnh (tính 1 ảnh)
            valid_list = [d for d in details if d.get("status") in ("valid", "duplicate")]
            warn_list = [d for d in details if d.get("status") in ("duplicate", "unrelated")]

            def build_valid_df(items):
                rows = []
                for idx, it in enumerate(items, 1):
                    cnt = it.get("count", 1)
                    st = it.get("status", "valid")
                    if st == "duplicate":
                        dup_cnt = max(1, cnt - 1)
                        st_label = f"Hợp lệ (Trùng {dup_cnt} ảnh)"
                        note = f"Tính 1 ảnh hợp lệ, cảnh báo có {dup_cnt} ảnh trùng"
                    else:
                        st_label = "Hợp lệ"
                        note = ""

                    rows.append({
                        "STT": idx,
                        "Ngày": it.get("date", "--"),
                        "Giờ": it.get("time", "--"),
                        "Nhân viên (SR)": it.get("sr", "--"),
                        "Công ty": it.get("company", "--"),
                        "Địa chỉ điểm bán": format_clean_address(it.get("address", "")),
                        "Mã Timemark": it.get("timemark_code", "--"),
                        "Trạng thái": st_label,
                        "Số ảnh chụp": cnt,
                        "Số ảnh tính": 1,
                        "Danh sách file ảnh": it.get("files", ""),
                        "Ghi chú": note
                    })
                return pd.DataFrame(rows)

            def build_warn_df(items):
                rows = []
                for idx, it in enumerate(items, 1):
                    cnt = it.get("count", 1)
                    st = it.get("status", "valid")
                    if st == "duplicate":
                        dup_cnt = max(1, cnt - 1)
                        st_label = f"Trùng {dup_cnt} ảnh (Đã tính 1)"
                        note = f"Phát hiện {dup_cnt} ảnh trùng lặp. Đã ghi nhận 1 ảnh vào sheet Điểm Bán Hợp Lệ."
                        dup_col = dup_cnt
                    else:
                        st_label = "Không liên quan"
                        note = it.get("reason", "Ảnh không hợp lệ")
                        dup_col = 0

                    rows.append({
                        "STT": idx,
                        "Ngày": it.get("date", "--"),
                        "Giờ": it.get("time", "--"),
                        "Nhân viên (SR)": it.get("sr", "--"),
                        "Công ty": it.get("company", "--"),
                        "Địa chỉ": format_clean_address(it.get("address", "")),
                        "Mã Timemark": it.get("timemark_code", "--"),
                        "Trạng thái": st_label,
                        "Số ảnh chụp": cnt,
                        "Số ảnh trùng": dup_col,
                        "Danh sách file ảnh": it.get("files", ""),
                        "Chi tiết cảnh báo": note
                    })
                return pd.DataFrame(rows)

            def build_kpi_df(items):
                staff_stats = {}
                for it in items:
                    if it.get("status") == "unrelated":
                        continue
                    sr = (it.get("sr") or "--").strip()
                    if sr not in staff_stats:
                        staff_stats[sr] = {
                            "total_stores": 0,
                            "standard_stores": 0,
                            "single_stores": 0,
                            "dup_stores": 0,
                            "total_photos": 0
                        }
                    staff_stats[sr]["total_stores"] += 1
                    c = it.get("count", 1)
                    staff_stats[sr]["total_photos"] += c
                    if c == 2:
                        staff_stats[sr]["standard_stores"] += 1
                    elif c > 2:
                        staff_stats[sr]["dup_stores"] += 1
                    else:
                        staff_stats[sr]["single_stores"] += 1

                rows = []
                for idx, (sr, s) in enumerate(staff_stats.items(), 1):
                    rate = f"{(s['standard_stores'] / s['total_stores'] * 100):.1f}%" if s['total_stores'] > 0 else "0%"
                    rows.append({
                        "STT": idx,
                        "Nhân viên (SR)": sr,
                        "Tổng điểm bán": s["total_stores"],
                        "Đạt chuẩn (2 ảnh)": s["standard_stores"],
                        "Thiếu ảnh (1 ảnh)": s["single_stores"],
                        "Chụp dư (>2 ảnh)": s["dup_stores"],
                        "Tổng số ảnh chụp": s["total_photos"],
                        "Tỷ lệ đạt chuẩn": rate
                    })
                return pd.DataFrame(rows)

            df_valid = build_valid_df(valid_list)
            df_kpi = build_kpi_df(details)
            df_warn = build_warn_df(warn_list)

            with pd.ExcelWriter(save_path, engine='openpyxl') as writer:
                df_valid.to_excel(writer, index=False, sheet_name="Điểm Bán Hợp Lệ")
                df_kpi.to_excel(writer, index=False, sheet_name="KPI Nhân Viên")
                df_warn.to_excel(writer, index=False, sheet_name="Trùng Lặp & Cảnh Báo")

                from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
                header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
                header_fill_blue = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
                header_fill_green = PatternFill(start_color="15803D", end_color="15803D", fill_type="solid")
                header_fill_amber = PatternFill(start_color="9A3412", end_color="9A3412", fill_type="solid")
                cell_font = Font(name="Segoe UI", size=10)
                thin_border = Border(
                    left=Side(style='thin', color='CBD5E1'),
                    right=Side(style='thin', color='CBD5E1'),
                    top=Side(style='thin', color='CBD5E1'),
                    bottom=Side(style='thin', color='CBD5E1')
                )

                sheet_configs = [
                    ("Điểm Bán Hợp Lệ", header_fill_blue),
                    ("KPI Nhân Viên", header_fill_green),
                    ("Trùng Lặp & Cảnh Báo", header_fill_amber)
                ]

                for sheet_name, fill in sheet_configs:
                    ws = writer.sheets[sheet_name]
                    for col_idx, col in enumerate(ws.columns, start=1):
                        max_len = 0
                        for r_idx, cell in enumerate(col, start=1):
                            cell.border = thin_border
                            if r_idx == 1:
                                cell.font = header_font
                                cell.fill = fill
                                cell.alignment = Alignment(horizontal="center", vertical="center")
                            else:
                                cell.font = cell_font
                                if col_idx in (1, 2, 3, 5, 7, 8, 9, 10):
                                    cell.alignment = Alignment(horizontal="center", vertical="center")
                                else:
                                    cell.alignment = Alignment(horizontal="left", vertical="center")

                            v_str = str(cell.value or "")
                            if len(v_str) > max_len:
                                max_len = len(v_str)

                        col_letter = ws.cell(row=1, column=col_idx).column_letter
                        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

            QMessageBox.information(self, "Thành công", f"Đã xuất file Excel thành công tại:\n{save_path}")
        except Exception as e:
            QMessageBox.critical(self, "Lỗi xuất Excel", f"Không thể lưu file:\n{str(e)}")

    def open_gemini_dialog(self, initial_tab: int = 0):
        engine = get_ocr_engine()
        raw_records = []

        if self.last_raw_ocr_items:
            raw_records = engine.format_gemini_export_records(self.last_raw_ocr_items)
        elif self.all_results and self.all_results.get("details"):
            # Lấy thông tin từ bảng details nếu không có raw_ocr_items
            details = self.all_results.get("details", [])
            for d in details:
                fl = d.get("file_list") or [d.get("files", "")]
                for f in fl:
                    f_name = os.path.basename(f) if f else ""
                    raw_records.append({
                        "id": len(raw_records) + 1,
                        "filename": f_name,
                        "sr": d.get("sr", "--"),
                        "company": d.get("company", "--"),
                        "address": d.get("address", ""),
                        "gps": d.get("gps", ""),
                        "date": d.get("date", "--/--/----"),
                        "time": d.get("time", "--:--"),
                        "raw_text": d.get("address", "")
                    })
        elif self.selected_files:
            # Chưa chạy OCR trước đó: thực hiện trích xuất nhanh các trường từ các ảnh đã chọn
            prog = QProgressDialog("Đang trích xuất dữ liệu nhanh từ ảnh để chuẩn bị xuất JSON...", "Hủy", 0, len(self.selected_files), self)
            prog.setWindowTitle("Trích xuất thông tin OCR")
            prog.setWindowModality(Qt.WindowModal)
            prog.show()
            QApplication.processEvents()

            quick_items = []
            for idx, path in enumerate(self.selected_files):
                if prog.wasCanceled():
                    break
                prog.setValue(idx)
                QApplication.processEvents()
                img = safe_read_cv2(path)
                if img is not None:
                    fn = os.path.basename(path)
                    info = engine.extract_timemark_info(img)
                    info["filename"] = fn
                    info["filepath"] = path
                    quick_items.append(info)
            prog.setValue(len(self.selected_files))
            if quick_items:
                self.last_raw_ocr_items = quick_items
                raw_records = engine.format_gemini_export_records(quick_items)

        # Mở hộp thoại Gemini Assistant Dialog
        dlg = GeminiAssistantDialog(
            raw_records=raw_records,
            on_apply=self.apply_gemini_cleaned_data,
            initial_tab=initial_tab if raw_records else 1,
            parent=self
        )
        dlg.exec()

    def open_web_gemini_evaluator(self):
        engine = get_ocr_engine()
        if not self.last_raw_ocr_items and not self.all_results:
            QMessageBox.warning(self, "Chưa có dữ liệu", "Vui lòng chọn ảnh và chạy phân tích OCR trước khi gửi cho Gemini!")
            return

        raw_records = []
        if self.last_raw_ocr_items:
            raw_records = engine.format_gemini_export_records(self.last_raw_ocr_items)
        elif self.all_results and "grouped_stores" in self.all_results:
            for store in self.all_results["grouped_stores"]:
                for it in store.get("items", []):
                    raw_records.append({
                        "id": len(raw_records) + 1,
                        "filename": it.get("filename", ""),
                        "sr": it.get("sr_name", "") or it.get("sr", ""),
                        "company": it.get("company_name", "") or it.get("company", ""),
                        "address": it.get("address", ""),
                        "gps": it.get("gps", ""),
                        "date": it.get("date", ""),
                        "time": it.get("time", "")
                    })

        prompt = engine.get_gemini_cleaning_prompt(raw_records)
        QApplication.clipboard().setText(prompt)

        # Mở trình duyệt web sang Gemini (kèm Extension nếu có Chrome)
        ext_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "extensions", "gemini_evaluator")
        launch_browser_with_extension("https://gemini.google.com/app", [ext_path])

        QMessageBox.information(
            self,
            "Đã mở Web Gemini",
            f"Đã sao chép dữ liệu JSON ({len(raw_records)} ảnh) và Prompt chấm điểm vào Clipboard!\n\n"
            f"Trang web Gemini đang được mở trên trình duyệt.\n"
            f"• Tiện ích Gemini Evaluator sẽ tự động nạp prompt và bấm gửi để chấm điểm bán.\n"
            f"• Hoặc bạn chỉ cần bấm Ctrl+V vào ô chat của Gemini để bắt đầu chấm điểm."
        )

    def apply_gemini_cleaned_data(self, cleaned_data: list):
        if not cleaned_data:
            return
        try:
            engine = get_ocr_engine()
            results = engine.group_from_cleaned_records(cleaned_data, existing_items=self.last_raw_ocr_items)
            self.on_worker_success(results)
            self.status_title.setText("Đã chấm điểm bán thành công từ dữ liệu làm sạch bởi Gemini")
            self.status_desc.setText(f"Dữ liệu được chuẩn hóa bởi Gemini ({len(cleaned_data)} ảnh, {results['summary']['valid_stores']} điểm bán).")
            QMessageBox.information(
                self,
                "Chấm điểm thành công",
                f"Đã đối soát và chấm điểm bán thành công từ dữ liệu làm sạch bởi Gemini!\n\n"
                f"• Tổng số ảnh: {len(cleaned_data)}\n"
                f"• Số điểm bán hợp lệ: {results['summary']['valid_stores']}\n"
                f"• Điểm bán đạt chuẩn (2 ảnh): {results['summary']['standard_stores']}\n"
                f"• Điểm thiếu / trùng: {results['summary']['warning_total']}\n\n"
                f"Dữ liệu địa chỉ, tên SR đã được cập nhật chuẩn đẹp vào bảng bên dưới!"
            )
        except Exception as e:
            QMessageBox.critical(self, "Lỗi chấm điểm", f"Không thể xử lý dữ liệu JSON đã làm sạch:\n{str(e)}")

    def resizeEvent(self, event):
        super().resizeEvent(event)
        if hasattr(self, 'current_preview_cv_img') and self.current_preview_cv_img is not None:
            self.adjust_preview_frame_aspect(self.current_preview_cv_img)
            pix = self.preview_img_label.pixmap()
            if pix and not pix.isNull():
                frame_w = max(self.preview_frame.width() - 8, 200)
                frame_h = max(self.preview_frame.height() - 8, 300)
                scaled = pix.scaled(frame_w, frame_h, Qt.KeepAspectRatio, Qt.SmoothTransformation)
                self.preview_img_label.setPixmap(scaled)


def main():
    app = QApplication(sys.argv)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()


